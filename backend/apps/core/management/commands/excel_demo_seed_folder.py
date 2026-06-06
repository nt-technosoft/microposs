"""Seed the local demo database from the investor Excel folder.

This command is intentionally a demo/workflow replay, not a report import:
it creates users/reference data directly, then runs procurements, receives,
transfers and sales through the same domain services used by the UI-facing
flows.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re

from django.apps import apps as django_apps
from django.contrib.auth.models import Group, User
from django.core.cache import cache
from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.demo_constants import DEFAULT_DEMO_USD_UZS_RATE
from apps.core.management.commands.excel_snapshot_from_xlsx import (
    DEFAULT_SHEETS,
    _is_required_row,
    _serialize,
)
from apps.core.management.commands.excel_workflow_audit import (
    Command as WorkflowEngine,
    _dec,
    _money,
    _parse_dt,
    _ratio,
    _unit,
)
from apps.core.management.safety import require_debug_or_confirmation
from apps.partnerships.formulas import profit_shares_from_capital


ROOT = Path(__file__).resolve().parents[5]
DEFAULT_FOLDER = ROOT / 'файлы'
MASTER_FILE = 'INVESTOR MASTER SHEET.xlsx'


@dataclass(frozen=True)
class DealSource:
    deal_id: str
    filename: str


DEALS = (
    DealSource('B-001', 'Copy of B-001 (1).xlsx'),
    DealSource('B-002', 'B-002 (1).xlsx'),
    DealSource('U-001', 'Copy of U-001 (1).xlsx'),
)


def _norm(text) -> str:
    return ' '.join(str(text or '').replace('\n', ' ').strip().upper().split())


def _product_key(value) -> str:
    return str(value or '').strip()


def _currency(value) -> str:
    raw = _norm(value)
    if raw in {'SOM', "SO'M"}:
        return 'UZS'
    return raw or 'UZS'


def _is_executable_sale(row: dict) -> bool:
    return bool(row.get('MAHSULOT')) and _dec(row.get('JAMI DONA')) > 0 and _dec(row.get('SOTUV NARXI')) > 0


def _username(label: str, used: set[str]) -> str:
    translit = {
        'А': 'a', 'Б': 'b', 'В': 'v', 'Г': 'g', 'Д': 'd', 'Е': 'e', 'Ё': 'e',
        'Ж': 'j', 'З': 'z', 'И': 'i', 'Й': 'y', 'К': 'k', 'Л': 'l', 'М': 'm',
        'Н': 'n', 'О': 'o', 'П': 'p', 'Р': 'r', 'С': 's', 'Т': 't', 'У': 'u',
        'Ф': 'f', 'Х': 'h', 'Ц': 'c', 'Ч': 'ch', 'Ш': 'sh', 'Щ': 'sh',
        'Ъ': '', 'Ы': 'y', 'Ь': '', 'Э': 'e', 'Ю': 'yu', 'Я': 'ya',
    }
    text = ''.join(translit.get(ch.upper(), ch.lower()) for ch in str(label))
    base = re.sub(r'[^a-z0-9]+', '_', text).strip('_') or 'user'
    candidate = base[:28]
    index = 2
    while candidate in used:
        suffix = f'_{index}'
        candidate = f'{base[:28 - len(suffix)]}{suffix}'
        index += 1
    used.add(candidate)
    return candidate


def _find_header(ws, first_label: str) -> int:
    target = _norm(first_label)
    for row in range(1, min(ws.max_row or 1, 30) + 1):
        if _norm(ws.cell(row, 1).value) == target:
            return row
    raise CommandError(f'Header {first_label!r} not found in {ws.title}.')


class Command(BaseCommand):
    help = 'Seed demo data from files/ investor Excel workbooks using UI-like domain flows.'

    def add_arguments(self, parser):
        parser.add_argument('--folder', default=str(DEFAULT_FOLDER))
        parser.add_argument('--apply', action='store_true')
        parser.add_argument('--dry-run', action='store_true')
        parser.add_argument('--wipe', action='store_true')
        parser.add_argument('--confirm-production-wipe', action='store_true')

    def handle(self, *args, **options):
        if bool(options['apply']) == bool(options['dry_run']):
            raise CommandError('Choose exactly one: --apply or --dry-run.')
        if options['apply'] and not options['wipe']:
            raise CommandError('--apply requires --wipe.')
        if options['apply']:
            require_debug_or_confirmation(
                command_name='excel_demo_seed_folder',
                confirmed=options['confirm_production_wipe'],
                flag_name='--confirm-production-wipe',
                action='this command truncates business data and rebuilds the Excel demo workflow',
            )

        folder = Path(options['folder']).expanduser()
        if not folder.exists():
            raise CommandError(f'Folder not found: {folder}')

        master = self._read_master(folder / MASTER_FILE)
        deal_books = {
            source.deal_id: self._with_opening_shortages(
                deal_id=source.deal_id,
                sheets=self._read_workbook(folder / source.filename),
            )
            for source in DEALS
        }
        plan = self._build_plan(master, deal_books)
        self._print_plan(plan)
        if options['dry_run']:
            return

        with transaction.atomic():
            self._wipe_database()
            ctx = self._create_baseline(master)
            for source in DEALS:
                self._apply_deal(
                    ctx=ctx,
                    deal_id=source.deal_id,
                    sheets=deal_books[source.deal_id],
                    funding_rows=master['funding_by_deal'].get(source.deal_id, []),
                )
        WorkflowEngine()._refresh_financial_aggregates()
        self.stdout.write(self.style.SUCCESS('Excel demo seed applied.'))
        self.stdout.write('Access:')
        self.stdout.write('  owner / Owner123!')
        self.stdout.write('  cashier / Cashier123!')
        self.stdout.write('  warehouse / Warehouse123!')
        self.stdout.write('  investor accounts use Investor123! password.')

    def _read_workbook(self, path: Path) -> dict:
        if not path.exists():
            raise CommandError(f'Workbook not found: {path}')
        wb = load_workbook(path, data_only=True)
        snapshot = {}
        for sheet_name in DEFAULT_SHEETS:
            if sheet_name not in wb.sheetnames:
                snapshot[sheet_name] = []
                continue
            ws = wb[sheet_name]
            headers = [
                str(cell.value).strip() if cell.value is not None else ''
                for cell in next(ws.iter_rows(min_row=1, max_row=1))
            ]
            rows = []
            for row_id, row_cells in enumerate(ws.iter_rows(min_row=2), start=2):
                row = {}
                for header, cell in zip(headers, row_cells):
                    if header:
                        row[header] = _serialize(cell.value)
                if all(value is None or str(value).strip() == '' for value in row.values()):
                    continue
                if not _is_required_row(sheet_name, row):
                    continue
                row['_row_id'] = str(row_id)
                rows.append(row)
            snapshot[sheet_name] = rows
        snapshot['_foyda'] = self._read_foyda(wb)
        return snapshot

    def _read_foyda(self, wb) -> list[dict]:
        if 'FOYDA_TAQSIMOTI' not in wb.sheetnames:
            raise CommandError('FOYDA_TAQSIMOTI sheet is required.')
        ws = wb['FOYDA_TAQSIMOTI']
        rows = []
        for row in range(2, min(ws.max_row or 2, 8) + 1):
            name = ws.cell(row, 1).value
            capital = ws.cell(row, 2).value
            capital_share = ws.cell(row, 3).value
            profit_share = ws.cell(row, 4).value
            try:
                capital_value = _money(capital)
                capital_share_value = _ratio(capital_share or 0)
                profit_share_value = _ratio(profit_share or 0)
            except Exception:
                continue
            if name and capital is not None and profit_share is not None:
                rows.append({
                    'name': str(name).strip(),
                    'capital': capital_value,
                    'capital_share': capital_share_value,
                    'profit_share': profit_share_value,
                })
        if len(rows) < 2:
            raise CommandError('FOYDA_TAQSIMOTI must contain investor/operator rows.')
        return rows

    def _read_master(self, path: Path) -> dict:
        if not path.exists():
            raise CommandError(f'Master workbook not found: {path}')
        wb = load_workbook(path, data_only=True)
        investors = []
        ws = wb['INVESTORLAR']
        header = _find_header(ws, 'Investor ID')
        for row in range(header + 1, (ws.max_row or 0) + 1):
            investor_id = ws.cell(row, 1).value
            name = ws.cell(row, 2).value
            status = ws.cell(row, 4).value
            if investor_id and name and _norm(status) == 'AKTIV':
                investors.append({'id': str(investor_id).strip(), 'name': str(name).strip()})

        funding_by_deal = defaultdict(list)
        ws = wb['BITIM_FUNDING']
        header = _find_header(ws, 'Sana')
        headers = [ws.cell(header, col).value for col in range(1, 9)]
        for row in range(header + 1, (ws.max_row or 0) + 1):
            record = {headers[col - 1]: ws.cell(row, col).value for col in range(1, len(headers) + 1)}
            deal_id = str(record.get('Bitim ID') or '').strip()
            investor = str(record.get('Investor') or '').strip()
            amount = record.get('Tikilgan summa USD')
            if deal_id and investor and amount not in (None, ''):
                funding_by_deal[deal_id].append({
                    'investor': investor,
                    'amount_usd': Decimal(str(amount)),
                    'date': record.get('Sana'),
                })
        return {'investors': investors, 'funding_by_deal': dict(funding_by_deal)}

    def _build_plan(self, master: dict, deal_books: dict) -> dict:
        result = {
            'investors': len(master['investors']),
            'deals': {},
        }
        for deal_id, sheets in deal_books.items():
            purchase_rows = sheets.get('SOTIB OLISH', [])
            sale_rows = [row for row in sheets.get('SOTUV', []) if _is_executable_sale(row)]
            currencies = Counter(_currency(row.get('PUL BIRLIGI')) for row in purchase_rows)
            purchase_total = sum(
                (_dec(row.get('SONI')) * _dec(row.get('MAHSULOT NARXI')) for row in purchase_rows),
                Decimal('0'),
            )
            result['deals'][deal_id] = {
                'purchase_rows': len(purchase_rows),
                'sale_rows': len(sale_rows),
                'purchase_total': _money(purchase_total),
                'currencies': dict(currencies),
                'foyda': sheets['_foyda'],
                'opening_shortages': sheets.get('_opening_shortages', []),
            }
        return result

    def _print_plan(self, plan: dict) -> None:
        self.stdout.write('Excel demo seed plan')
        self.stdout.write(f"  investor accounts: {plan['investors']}")
        for deal_id, row in plan['deals'].items():
            self.stdout.write(
                f"  {deal_id}: purchases={row['purchase_rows']}, sales={row['sale_rows']}, "
                f"purchase_total={row['purchase_total']}, currencies={row['currencies']}"
            )
            for partner in row['foyda']:
                self.stdout.write(
                    f"    {partner['name']}: capital={partner['capital']}, "
                    f"capital_share={partner['capital_share']}, profit_share={partner['profit_share']}"
                )
            for shortage in row['opening_shortages']:
                self.stdout.write(
                    f"    opening stock: {shortage['product']} +{shortage['quantity']} "
                    f"at {shortage['unit_cost']} {shortage['currency']}"
                )

    def _with_opening_shortages(self, *, deal_id: str, sheets: dict) -> dict:
        purchase_rows = [row for row in sheets.get('SOTIB OLISH', []) if row.get('MAHSULOT')]
        sale_rows = [row for row in sheets.get('SOTUV', []) if _is_executable_sale(row)]
        purchase_qty = defaultdict(Decimal)
        purchase_cost = defaultdict(Decimal)
        purchase_canonical = {}
        purchase_templates = {}
        for row in purchase_rows:
            raw = _product_key(row.get('MAHSULOT'))
            key = _norm(raw)
            qty = _dec(row.get('SONI'))
            unit_cost = _dec(row.get('MAHSULOT NARXI'))
            purchase_qty[key] += qty
            purchase_cost[key] += qty * unit_cost
            purchase_canonical.setdefault(key, raw)
            purchase_templates.setdefault(key, row)

        sale_qty = defaultdict(Decimal)
        for row in sale_rows:
            sale_qty[_norm(_product_key(row.get('MAHSULOT')))] += _dec(row.get('JAMI DONA'))

        synthetic_rows = []
        opening_shortages = []
        for key, sold_qty in sorted(sale_qty.items()):
            bought_qty = purchase_qty.get(key, Decimal('0'))
            if sold_qty <= bought_qty:
                continue
            if key not in purchase_templates:
                raise CommandError(f'{deal_id}: sales reference unknown product {_product_key(key)}.')
            shortage = sold_qty - bought_qty
            template = dict(purchase_templates[key])
            unit_cost = _unit(purchase_cost[key] / bought_qty) if bought_qty > 0 else _unit(template.get('MAHSULOT NARXI'))
            template['MAHSULOT'] = purchase_canonical[key]
            template['SONI'] = shortage
            template['MAHSULOT NARXI'] = unit_cost
            template['_row_id'] = f'opening-shortage-{deal_id}-{purchase_canonical[key]}'
            template['_demo_note'] = 'Opening stock reconstructed because Excel sales exceed purchase quantity.'
            synthetic_rows.append(template)
            opening_shortages.append({
                'product': purchase_canonical[key],
                'quantity': shortage,
                'unit_cost': unit_cost,
                'currency': _currency(template.get('PUL BIRLIGI')),
            })

        if not synthetic_rows:
            return sheets
        return {
            **sheets,
            'SOTIB OLISH': [*purchase_rows, *synthetic_rows],
            '_opening_shortages': opening_shortages,
        }

    def _wipe_database(self) -> None:
        cache.clear()
        keep = {'auth', 'admin', 'contenttypes', 'sessions', 'django_celery_beat'}
        tables = [
            model._meta.db_table
            for model in django_apps.get_models()
            if model._meta.app_label not in keep
        ]
        if tables:
            quoted = ', '.join(f'"{table}"' for table in tables)
            with connection.cursor() as cursor:
                cursor.execute(f'TRUNCATE {quoted} RESTART IDENTITY CASCADE;')
        User.objects.all().delete()

    def _create_baseline(self, master: dict) -> dict:
        from apps.catalog.models import Category, DiscountReason
        from apps.core.models import Business, BusinessInvestorRelation, Partner
        from apps.finance.chart_of_accounts import setup_chart_of_accounts
        from apps.finance.models import Account, CashAccount, ExchangeRate
        from apps.inventory.models import Warehouse
        from apps.investors.models import Investor
        from apps.suppliers.models import Supplier

        groups = {name: Group.objects.get_or_create(name=name)[0] for name in ('owner', 'cashier', 'warehouse', 'investor')}
        used = set()

        def create_user(username, email, password, group, *, superuser=False):
            user = User.objects.create_user(username=username, email=email, password=password)
            user.is_active = True
            user.is_staff = superuser
            user.is_superuser = superuser
            user.save(update_fields=['is_active', 'is_staff', 'is_superuser'])
            user.groups.set([group])
            used.add(username)
            return user

        owner = create_user('owner', 'owner@local.dev', 'Owner123!', groups['owner'], superuser=True)
        cashier = create_user('cashier', 'cashier@local.dev', 'Cashier123!', groups['cashier'])
        warehouse_user = create_user('warehouse', 'warehouse@local.dev', 'Warehouse123!', groups['warehouse'])
        business = Business.objects.create(owner=owner, name='Sherik Demo Excel', currency='UZS', is_active=True)
        tenant_id = business.id
        setup_chart_of_accounts(tenant_id)
        ExchangeRate.objects.create(
            tenant=business,
            base_currency='USD',
            quote_currency='UZS',
            rate_date=timezone.localdate(),
            rate=DEFAULT_DEMO_USD_UZS_RATE,
            source=ExchangeRate.Source.MANUAL,
            is_manual=True,
            notes='Excel demo seed fallback rate',
            raw_payload={},
            fetched_at=timezone.now(),
        )

        cash_account = Account.objects.get(tenant_id=tenant_id, code='1000')
        bank_account = Account.objects.get(tenant_id=tenant_id, code='1010')
        cash_accounts = {}
        for name, currency, kind, linked in (
            ('KASSA SOM', 'UZS', CashAccount.Kind.CASH, cash_account),
            ('KASSA DOLLAR', 'USD', CashAccount.Kind.CASH, bank_account),
            ('PLASTIK SOM', 'UZS', CashAccount.Kind.CARD_TERMINAL, bank_account),
        ):
            cash_accounts[name] = CashAccount.objects.create(
                tenant=business,
                name=name,
                currency=currency,
                kind=kind,
                balance=Decimal('0'),
                is_active=True,
                linked_account=linked,
            )
        warehouses = {
            'ASOSIY': Warehouse.objects.create(tenant=business, name='Основной склад', kind=Warehouse.WarehouseKind.STORAGE, is_active=True),
            'DOKON': Warehouse.objects.create(tenant=business, name='Основной магазин', kind=Warehouse.WarehouseKind.SHOP, is_active=True),
        }
        operator = Partner.objects.create(tenant=business, role=Partner.Role.OPERATOR, display_name='Owner operator', user=owner, is_active=True)

        investor_partners = {}
        investor_users = {}
        for investor in master['investors']:
            username = _username(investor['name'], used)
            user = create_user(username, f'{username}@local.dev', 'Investor123!', groups['investor'])
            partner = Partner.objects.create(
                tenant=business,
                role=Partner.Role.INVESTOR,
                display_name=investor['name'],
                user=user,
                is_active=True,
            )
            BusinessInvestorRelation.objects.create(
                tenant=business,
                partner=partner,
                status=BusinessInvestorRelation.Status.ACTIVE,
                source=BusinessInvestorRelation.Source.MANUAL,
                created_by=owner,
                notes='Excel demo investor access',
            )
            Investor.objects.create(tenant=business, user=user, name=investor['name'], email=user.email, is_active=True)
            investor_partners[investor['name']] = partner
            investor_users[investor['name']] = user

        category = Category.objects.create(tenant=business, name='Excel demo товары', sort_order=1)
        discount_reason = DiscountReason.objects.create(tenant=business, name='Excel price', is_default=True, is_active=True)
        supplier = Supplier.objects.create(tenant=business, name='Excel supplier', is_active=True)
        return {
            'business': business,
            'tenant_id': tenant_id,
            'owner': owner,
            'cashier': cashier,
            'warehouse_user': warehouse_user,
            'operator': operator,
            'investor_partners': investor_partners,
            'category': category,
            'discount_reason': discount_reason,
            'supplier': supplier,
            'warehouses': warehouses,
            'cash_accounts': cash_accounts,
        }

    def _apply_deal(self, *, ctx: dict, deal_id: str, sheets: dict, funding_rows: list[dict]) -> None:
        from apps.customers.models import Customer

        customers = {}
        for name in sorted({str(row.get('MIJOZ') or '').strip() for row in sheets.get('SOTUV', []) if str(row.get('MIJOZ') or '').strip()}):
            customers[name] = Customer.objects.get_or_create(
                tenant_id=ctx['tenant_id'],
                name=name,
                defaults={'is_active': True},
            )[0]
        variants = self._create_products(ctx=ctx, deal_id=deal_id, sheets=sheets)
        self._seed_fx_rates(ctx=ctx, deal_id=deal_id, sheets=sheets)
        plan = WorkflowEngine()._build_plan(sheets)
        procurement = self._create_procurement(
            ctx=ctx,
            deal_id=deal_id,
            sheets=sheets,
            plan=plan,
            variants=variants,
            funding_rows=funding_rows,
        )
        WorkflowEngine()._transfer_for_sales(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            warehouses=ctx['warehouses'],
            sheets=sheets,
            plan=plan,
            variants=variants,
        )
        sale_sheets = {**sheets, 'SOTUV': [row for row in sheets.get('SOTUV', []) if _is_executable_sale(row)]}
        WorkflowEngine()._replay_sales(
            tenant_id=ctx['tenant_id'],
            cashier=ctx['cashier'],
            shop=ctx['warehouses']['DOKON'],
            cash_accounts=ctx['cash_accounts'],
            customers=customers,
            variants=variants,
            discount_reason_id=ctx['discount_reason'].id,
            sheets=sale_sheets,
        )
        self.stdout.write(self.style.SUCCESS(f'  {deal_id}: applied'))

    def _create_products(self, *, ctx: dict, deal_id: str, sheets: dict) -> dict:
        from apps.catalog.services import create_product_with_variants

        products = {_product_key(row['MAHSULOT']) for row in sheets.get('SOTIB OLISH', []) if row.get('MAHSULOT')}
        sale_price_modes = defaultdict(Counter)
        for row in sheets.get('SOTUV', []):
            raw_name = _product_key(row.get('MAHSULOT'))
            if not raw_name:
                continue
            currency = _currency(row.get('VALYUTA'))
            raw = _dec(row.get('SOTUV NARXI'))
            fx = _dec(row.get('KURS'), '1')
            price_uzs = _money(raw if currency == 'UZS' else raw * fx)
            sale_price_modes[raw_name][price_uzs] += 1
        variants = {}
        normalized_variants = {}
        for raw_name in sorted(products):
            mode = sale_price_modes[raw_name].most_common(1)
            base_price = mode[0][0] if mode else Decimal('100000.00')
            product = create_product_with_variants(
                tenant_id=ctx['tenant_id'],
                name=f'{deal_id} · {raw_name}',
                category_id=ctx['category'].id,
                base_price=str(base_price),
                pricing_mode='EDITABLE',
                variant_data=None,
            )
            variant = product.variants.filter(is_active=True).first()
            variants[raw_name] = variant
            normalized_variants[_norm(raw_name)] = variant
        for sheet_name in ('SOTUV', 'STOCK TRANSFER'):
            for row in sheets.get(sheet_name, []):
                raw_name = _product_key(row.get('MAHSULOT'))
                if raw_name and raw_name not in variants and _norm(raw_name) in normalized_variants:
                    variants[raw_name] = normalized_variants[_norm(raw_name)]
        return variants

    def _seed_fx_rates(self, *, ctx: dict, deal_id: str, sheets: dict) -> None:
        from apps.finance.models import ExchangeRate

        rates_by_date = {}
        for row in sheets.get('SOTIB OLISH', []):
            if _currency(row.get('PUL BIRLIGI')) != 'USD':
                continue
            date_value = (
                row.get('MAHSULOT OMBORGA YETIB KELGAN SANA')
                or row.get("MAHSULOT DO'KONGA YETIB KELGAN SANA")
                or row.get('SANA')
            )
            if not date_value:
                continue
            rates_by_date[_parse_dt(date_value).date()] = _dec(row.get('KURS'), DEFAULT_DEMO_USD_UZS_RATE)

        purchase_currency = None
        purchase_rows = sheets.get('SOTIB OLISH', [])
        if purchase_rows:
            purchase_currency = _currency(purchase_rows[0].get('PUL BIRLIGI'))
        for row in sheets.get('SOTUV', []):
            date_value = row.get('SOTUV SANASI')
            if not date_value:
                continue
            sale_currency = _currency(row.get('VALYUTA'))
            if sale_currency != 'USD' and purchase_currency != 'USD':
                continue
            rates_by_date[_parse_dt(date_value).date()] = _dec(row.get('KURS'), DEFAULT_DEMO_USD_UZS_RATE)

        for rate_date, rate in rates_by_date.items():
            ExchangeRate.objects.update_or_create(
                tenant_id=ctx['tenant_id'],
                base_currency='USD',
                quote_currency='UZS',
                rate_date=rate_date,
                defaults={
                    'rate': rate,
                    'source': ExchangeRate.Source.MANUAL,
                    'is_manual': True,
                    'notes': f'Excel demo FX {deal_id}',
                    'raw_payload': {'source': deal_id},
                    'fetched_at': timezone.now(),
                },
            )

    def _create_procurement(self, *, ctx: dict, deal_id: str, sheets: dict, plan: dict, variants: dict, funding_rows: list[dict]):
        from apps.inventory.models import Warehouse
        from apps.partnerships.workspace import create_workspace, dispatch_workspace_action

        purchase_rows = plan['purchase_rows']
        if not purchase_rows:
            raise CommandError(f'{deal_id}: no purchase rows.')
        purchase_currency = _currency(purchase_rows[0].get('PUL BIRLIGI'))
        if any(_currency(row.get('PUL BIRLIGI')) != purchase_currency for row in purchase_rows):
            raise CommandError(f'{deal_id}: mixed purchase currencies are not supported by this demo seed.')
        purchase_total = _money(sum((_dec(row.get('SONI')) * _dec(row.get('MAHSULOT NARXI')) for row in purchase_rows), Decimal('0')))
        purchase_fx = _dec(purchase_rows[0].get('KURS'), DEFAULT_DEMO_USD_UZS_RATE) if purchase_currency == 'USD' else Decimal('1')
        received_at = _parse_dt(
            purchase_rows[0].get('MAHSULOT OMBORGA YETIB KELGAN SANA')
            or purchase_rows[0].get("MAHSULOT DO'KONGA YETIB KELGAN SANA")
            or purchase_rows[0].get('SANA')
        )

        partners_payload, allocations = self._agreement_payload(
            ctx=ctx,
            deal_id=deal_id,
            currency=purchase_currency,
            purchase_total=purchase_total,
            foyda_rows=sheets['_foyda'],
            funding_rows=funding_rows,
            fx_rate=purchase_fx,
        )
        mudaraba_ratio = self._mudaraba_ratio(partners_payload)

        procurement = create_workspace(
            tenant_id=ctx['tenant_id'],
            funding_source='PARTNERSHIP',
            primary_currency=purchase_currency,
            supplier_id=ctx['supplier'].id,
            notes=f'Excel demo {deal_id}',
        )
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='CREATE_INVESTMENT_AGREEMENT',
            payload={'payload': {
                'supplier_id': ctx['supplier'].id,
                'mudaraba_ratio': str(mudaraba_ratio),
                'planned_budget': str(purchase_total),
                'currency': purchase_currency,
                'reconciliation_mode': 'FACTUAL',
                'default_advance_repayment_mode': 'LUMP',
                'notes': f'Excel demo agreement {deal_id}',
                'partners': partners_payload,
            }},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        for allocation in allocations:
            procurement = dispatch_workspace_action(
                tenant_id=ctx['tenant_id'],
                procurement=procurement,
                action='RECORD_CAPITAL_CONTRIBUTION',
                payload={'payload': {
                    'partner_id': allocation['partner_id'],
                    'amount': allocation['amount'],
                    'currency': purchase_currency,
                    'fx_rate': allocation['fx_rate'],
                    'date': received_at.date().isoformat(),
                    'notes': f'Excel demo capital {deal_id}',
                }},
                user_id=ctx['owner'].id,
            )

        item_payloads = [{
            'product_variant_id': variants[_product_key(row['MAHSULOT'])].id,
            'quantity': str(_dec(row.get('SONI'))),
            'unit_purchase_price': str(_unit(row.get('MAHSULOT NARXI'))),
            'currency': purchase_currency,
            'fx_rate': str(_dec(row.get('KURS'), DEFAULT_DEMO_USD_UZS_RATE) if purchase_currency == 'USD' else Decimal('1')),
        } for row in purchase_rows]
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='UPDATE_ITEMS',
            payload={'payload': {'items': item_payloads}},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        item_ids = list(procurement.items.order_by('id').values_list('id', flat=True))
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='UPDATE_SETTLEMENT',
            payload={'payload': {
                'type': 'PREPAID',
                'currency_of_obligation': purchase_currency,
                'fx_rate_at_obligation': str(purchase_fx),
                'total_amount_due': str(purchase_total),
                'notes': f'Excel demo settlement {deal_id}',
            }},
            user_id=ctx['owner'].id,
        )
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='ALLOCATE_CAPITAL',
            payload={'payload': {
                'item_ids': item_ids,
                'expense_ids': [],
                'allocations': allocations,
                'date': received_at.date().isoformat(),
                'notes': f'Excel demo allocation {deal_id}',
            }},
            user_id=ctx['owner'].id,
        )
        storage = Warehouse.objects.get(tenant_id=ctx['tenant_id'], name='Основной склад')
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='RECEIVE_BATCH',
            payload={'payload': {
                'warehouse_id': storage.id,
                'item_ids': item_ids,
                'capital_allocations': allocations,
                'received_at': received_at.isoformat(),
            }},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        return procurement

    def _agreement_payload(self, *, ctx: dict, deal_id: str, currency: str, purchase_total: Decimal, foyda_rows: list[dict], funding_rows: list[dict], fx_rate: Decimal):
        investor_aggregate = next((row for row in foyda_rows if 'USTOZ' in _norm(row['name'])), foyda_rows[0])
        operator_row = next((row for row in foyda_rows if row is not investor_aggregate), foyda_rows[-1])

        aggregate_capital = _money(investor_aggregate['capital'])
        if aggregate_capital <= 0:
            aggregate_capital = purchase_total
        if funding_rows:
            total_usd = sum((Decimal(str(row['amount_usd'])) for row in funding_rows), Decimal('0'))
            funding_amounts = defaultdict(Decimal)
            for row in funding_rows:
                funding_amounts[row['investor']] += Decimal(str(row['amount_usd']))
            funding = [
                {'investor': investor, 'amount': _money(aggregate_capital * amount / total_usd)}
                for investor, amount in funding_amounts.items()
                if total_usd > 0
            ]
        else:
            funding = [{'investor': investor_aggregate['name'], 'amount': aggregate_capital}]

        operator_capital = _money(max(purchase_total - sum((row['amount'] for row in funding), Decimal('0')), Decimal('0')))
        aggregate_profit_share = Decimal(str(investor_aggregate['profit_share']))
        partners_payload = []
        investor_profit_sum = Decimal('0')
        for row in funding:
            partner = ctx['investor_partners'].get(row['investor'])
            if partner is None:
                partner = self._create_ad_hoc_investor(ctx, row['investor'])
            share = _ratio(aggregate_profit_share * row['amount'] / sum((item['amount'] for item in funding), Decimal('0')))
            investor_profit_sum += share
            partners_payload.append({
                'partner_id': partner.id,
                'role': 'INVESTOR',
                'planned_capital_share': str(row['amount']),
                'profit_share': str(share),
            })
        partners_payload.append({
            'partner_id': ctx['operator'].id,
            'role': 'OPERATOR',
            'planned_capital_share': str(operator_capital),
            'profit_share': str(_ratio(Decimal('1') - investor_profit_sum)),
        })

        # Recompute profit shares from the contract formula so validator and seed agree.
        mudaraba_ratio = self._mudaraba_ratio(partners_payload)
        meta = []
        for partner in partners_payload:
            meta.append({
                'partner_id': partner['partner_id'],
                'role': partner['role'],
                'capital_share': _ratio(Decimal(str(partner['planned_capital_share'])) / purchase_total),
            })
        expected = profit_shares_from_capital(meta, mudaraba_ratio)
        assigned_profit = Decimal('0')
        for index, partner in enumerate(partners_payload):
            share = _ratio(expected[int(partner['partner_id'])])
            if index == len(partners_payload) - 1:
                share = _ratio(Decimal('1') - assigned_profit)
            assigned_profit += share
            partner['profit_share'] = str(share)

        allocations = []
        allocated = Decimal('0')
        for index, partner in enumerate(partners_payload):
            amount = _money(partner['planned_capital_share'])
            if index == len(partners_payload) - 1:
                amount = _money(purchase_total - allocated)
            allocated += amount
            allocations.append({
                'partner_id': partner['partner_id'],
                'amount': str(amount),
                'currency': currency,
                'fx_rate': str(fx_rate if currency == 'USD' else Decimal('1')),
            })
        return partners_payload, allocations

    def _mudaraba_ratio(self, partners_payload: list[dict]) -> Decimal:
        total = sum((Decimal(str(partner['planned_capital_share'])) for partner in partners_payload), Decimal('0'))
        investor_capital = sum(
            (Decimal(str(partner['planned_capital_share'])) for partner in partners_payload if partner['role'] == 'INVESTOR'),
            Decimal('0'),
        )
        investor_profit = sum(
            (Decimal(str(partner['profit_share'])) for partner in partners_payload if partner['role'] == 'INVESTOR'),
            Decimal('0'),
        )
        if total <= 0 or investor_capital <= 0:
            return Decimal('0')
        return _ratio(investor_profit / (investor_capital / total))

    def _create_ad_hoc_investor(self, ctx: dict, name: str):
        from apps.core.models import BusinessInvestorRelation, Partner
        from apps.investors.models import Investor

        group = Group.objects.get(name='investor')
        used = set(User.objects.values_list('username', flat=True))
        username = _username(name, used)
        user = User.objects.create_user(username=username, email=f'{username}@local.dev', password='Investor123!')
        user.is_active = True
        user.save(update_fields=['is_active'])
        user.groups.set([group])
        partner = Partner.objects.create(tenant=ctx['business'], role=Partner.Role.INVESTOR, display_name=name, user=user, is_active=True)
        BusinessInvestorRelation.objects.create(
            tenant=ctx['business'],
            partner=partner,
            status=BusinessInvestorRelation.Status.ACTIVE,
            source=BusinessInvestorRelation.Source.MANUAL,
            created_by=ctx['owner'],
            notes='Excel demo ad-hoc investor access',
        )
        Investor.objects.create(tenant=ctx['business'], user=user, name=name, email=user.email, is_active=True)
        ctx['investor_partners'][name] = partner
        return partner
