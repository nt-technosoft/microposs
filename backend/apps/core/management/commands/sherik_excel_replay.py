"""Preflight and apply Sherik Excel Replay through user-facing domain services.

The command always validates source workbooks first. Apply mode writes a backup,
wipes only after a clean preflight, then replays onboarding, agreements,
procurements, receipts, transfers and sales through domain service entrypoints.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
import re
import secrets
import string
from uuid import NAMESPACE_URL, uuid5

from django.apps import apps as django_apps
from django.contrib.auth.models import Group, User
from django.core.cache import cache
from django.core.management import call_command
from django.core.management.base import BaseCommand, CommandError
from django.db import connection, transaction
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.management.safety import require_debug_or_confirmation


ROOT = Path(__file__).resolve().parents[5]
DEFAULT_FOLDER = ROOT / 'файлы'
MASTER_FILE = 'INVESTOR MASTER SHEET.xlsx'

ZERO = Decimal('0')
CENT = Decimal('0.01')
RATIO = Decimal('0.000001')


@dataclass(frozen=True)
class AgreementTerm:
    name: str
    capital: Decimal
    capital_share: Decimal
    profit_share: Decimal


@dataclass
class LocalDeal:
    deal_id: str
    path: Path
    terms_currency: str
    terms: list[AgreementTerm]
    rows: dict[str, list[dict]]
    source_corrections: list[str] = field(default_factory=list)


SOURCE_CORRECTIONS = {
    'D-001': {
        'purchase_received_at_overrides': [
            {
                'row_id': 15,
                'product': 'DETSKIY GORKA 2-2',
                'received_at': '2026-07-03',
                'reason': (
                    'post-factum contract-2 Excel entry without warehouse arrival date; '
                    'E23 rollover requires recovered capital from completed contract-1 sales'
                ),
            },
            {
                'row_id': 16,
                'product': 'KRESLO-2',
                'received_at': '2026-07-03',
                'reason': (
                    'post-factum contract-2 Excel entry without warehouse arrival date; '
                    'E23 rollover requires recovered capital from completed contract-1 sales'
                ),
            },
            {
                'row_id': 17,
                'product': 'TRENAJOR-2',
                'received_at': '2026-07-03',
                'reason': (
                    'post-factum contract-2 Excel entry without warehouse arrival date; '
                    'E23 rollover requires recovered capital from completed contract-1 sales'
                ),
            },
            {
                'row_id': 18,
                'product': 'DETSKIY GORKA 1-2',
                'received_at': '2026-07-03',
                'reason': (
                    'post-factum contract-2 Excel entry without warehouse arrival date; '
                    'E23 rollover requires recovered capital from completed contract-1 sales'
                ),
            },
        ],
    },
    'U-001': {
        'funding_date_overrides': [
            {
                'row_id': 9,
                'investor': 'Alijon Ravshanov',
                'effective_date': '2026-03-15',
                'reason': (
                    'founder-approved post-factum funding correction: March 23 new-money '
                    'top-up was available before the March 15 receipt shortfall'
                ),
            },
            {
                'row_id': 10,
                'investor': 'Azizbek Norqarayev',
                'effective_date': '2026-03-15',
                'reason': (
                    'founder-approved post-factum funding correction: March 23 new-money '
                    'top-up was available before the March 15 receipt shortfall'
                ),
            },
        ],
        'sale_product_reclassifications': [
            {
                'row_id': 488,
                'from': 'J82350',
                'to': '82350',
                'reason': 'founder-approved SKU split: keep J82350 at 20 sold and move one sale to 82350',
            },
        ],
        'purchase_received_at_overrides': [
            {
                'row_id': 29,
                'product': 'KURTKA 8803',
                'received_at': '2025-11-07',
                'reason': 'post-factum Excel entry; first sale predates recorded purchase',
            },
            {
                'row_id': 135,
                'product': '99620',
                'received_at': '2026-05-22',
                'reason': 'post-factum Excel entry; first sale is one day before recorded purchase',
            },
        ],
    },
}


def norm(value) -> str:
    return ' '.join(str(value or '').replace('\n', ' ').strip().upper().split())


def display_name(value) -> str:
    return ' '.join(str(value or '').replace('\n', ' ').strip().split())


def is_blank(value) -> bool:
    return value is None or str(value).strip() == ''


def dec(value, default: Decimal = ZERO) -> Decimal:
    if value in (None, ''):
        return default
    if isinstance(value, Decimal):
        return value
    if isinstance(value, str):
        text = value.strip().replace('\xa0', '').replace(' ', '').replace(',', '.')
        if text.endswith('%'):
            text = text[:-1]
        if not text:
            return default
        try:
            return Decimal(text)
        except InvalidOperation:
            return default
    try:
        return Decimal(str(value))
    except Exception:
        return default


def money(value) -> Decimal:
    return dec(value).quantize(CENT, rounding=ROUND_HALF_UP)


def ratio(value) -> Decimal:
    raw = dec(value)
    if raw > Decimal('1') and raw <= Decimal('100'):
        raw = raw / Decimal('100')
    return raw.quantize(RATIO, rounding=ROUND_HALF_UP)


def fmt(value: Decimal | int | float | str) -> str:
    amount = dec(value)
    text = f'{amount:,.6f}'.rstrip('0').rstrip('.')
    return text.replace(',', ' ')


def currency(value) -> str:
    raw = norm(value)
    if raw in {'SOM', "SO'M", 'SUM'}:
        return 'UZS'
    if raw in {'USD', 'DOLLAR', '$'}:
        return 'USD'
    return raw or 'UZS'


def date_label(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return str(value)[:10] if value else ''


def parse_dt(value):
    if value is None or value == '':
        return timezone.now()
    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).replace('T', ' ').split('.')[0]
        try:
            dt = datetime.strptime(text, '%Y-%m-%d %H:%M:%S')
        except ValueError:
            try:
                dt = datetime.strptime(text[:10], '%Y-%m-%d')
            except ValueError:
                return timezone.now()
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def pick(row: dict, *keys: str):
    targets = {norm(key) for key in keys}
    for key, value in row.items():
        if norm(key) in targets:
            return value
    return ''


def product_key(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r'\s+', ' ', str(value or '').strip().upper())


def apply_source_corrections(deal_id: str, rows: dict[str, list[dict]]) -> list[str]:
    corrections = SOURCE_CORRECTIONS.get(deal_id, {})
    applied: list[str] = []

    sale_rows = {
        int(row.get('_row_id') or 0): row
        for row in rows.get('SOTUV', [])
    }
    for spec in corrections.get('sale_product_reclassifications', []):
        row = sale_rows.get(spec['row_id'])
        if row is None:
            raise CommandError(f"{deal_id}: source correction sale row {spec['row_id']} was not found.")
        current = product_key(pick(row, 'MAHSULOT'))
        expected = product_key(spec['from'])
        if current != expected:
            raise CommandError(
                f"{deal_id}: source correction sale row {spec['row_id']} expected {expected}, got {current}."
            )
        row['MAHSULOT'] = spec['to']
        applied.append(
            f"SOTUV row {spec['row_id']}: product {spec['from']} -> {spec['to']} ({spec['reason']})"
        )

    purchase_rows = {
        int(row.get('_row_id') or 0): row
        for row in rows.get('SOTIB OLISH', [])
    }
    for spec in corrections.get('purchase_received_at_overrides', []):
        row = purchase_rows.get(spec['row_id'])
        if row is None:
            raise CommandError(f"{deal_id}: source correction purchase row {spec['row_id']} was not found.")
        current = product_key(pick(row, 'MAHSULOT'))
        expected = product_key(spec['product'])
        if current != expected:
            raise CommandError(
                f"{deal_id}: source correction purchase row {spec['row_id']} expected {expected}, got {current}."
            )
        row['_effective_received_at'] = spec['received_at']
        applied.append(
            f"SOTIB OLISH row {spec['row_id']}: effective received_at={spec['received_at']} "
            f"for {spec['product']} ({spec['reason']})"
        )

    credit_deferrals = apply_credit_term_receive_deferrals(rows)
    if credit_deferrals:
        applied.extend(credit_deferrals)

    return applied


def apply_credit_term_receive_deferrals(rows: dict[str, list[dict]]) -> list[str]:
    sale_dates: dict[str, list] = defaultdict(list)
    all_sale_dates = []
    for row in rows.get('SOTUV', []):
        product = product_key(pick(row, 'MAHSULOT'))
        if not product:
            continue
        sold_at = parse_dt(pick(row, 'SOTUV SANASI'))
        sale_dates[product].append(sold_at)
        all_sale_dates.append(sold_at)
    if not all_sale_dates:
        return []

    first_sale_by_product = {
        product: min(dates)
        for product, dates in sale_dates.items()
        if dates
    }
    after_last_sale = max(all_sale_dates) + timedelta(days=1)
    changed = []
    for row in rows.get('SOTIB OLISH', []):
        if row.get('_effective_received_at'):
            continue
        term = norm(pick(row, "TO'LOV MUDDATI"))
        if term not in {'NASIYA', 'COD'}:
            continue
        product = product_key(pick(row, 'MAHSULOT'))
        if not product:
            continue
        current = parse_dt(
            pick(
                row,
                'SANA',
                'MAHSULOT OMBORGA YETIB KELGAN SANA',
                "MAHSULOT DO'KONGA YETIB KELGAN SANA",
            )
        )
        target = first_sale_by_product.get(product) or after_last_sale
        if target <= current:
            continue
        effective = target.date().isoformat()
        row['_effective_received_at'] = effective
        changed.append(f"{int(row.get('_row_id') or 0)}->{effective}")

    if not changed:
        return []
    sample = ', '.join(changed[:16])
    suffix = ' ...' if len(changed) > 16 else ''
    return [
        'SOTIB OLISH COD/NASIYA effective capital-funded received_at deferred for '
        f'{len(changed)} row(s): {sample}{suffix}'
    ]


def business_key(deal_id: str, operator_name: str) -> str:
    prefix = (deal_id or '')[:1].upper()
    if prefix == 'B':
        return 'Bekzod'
    if prefix == 'D':
        return 'Doniyor'
    if prefix == 'U':
        return 'Uygun'
    return operator_name or prefix or 'Unknown'


def clean_person_name(value: str) -> str:
    name = display_name(value)
    name = re.sub(r'\b(aka|акя|ака)\b\.?', '', name, flags=re.IGNORECASE)
    return ' '.join(name.split())


def slug(value: str, *, prefix: str = '') -> str:
    base = re.sub(r'[^a-z0-9]+', '_', str(value or '').lower()).strip('_')
    translit = {
        'bekzod': 'bekzod',
        'doniyor': 'doniyor',
        'uygun': 'uygun',
        'alijon': 'alijon',
        'azizbek': 'azizbek',
        'kamoliddin': 'kamoliddin',
        'faxriyor': 'faxriyor',
        'nilufar': 'nilufar',
        'abdulloh': 'abdulloh',
        'nurbol': 'nurbol',
        'muhammad': 'muhammad',
        'muhammadumar': 'muhammadumar',
    }
    if not base:
        ascii_bits = []
        for part in re.split(r'\s+', str(value or '').lower()):
            ascii_bits.append(translit.get(part, ''))
        base = '_'.join(bit for bit in ascii_bits if bit) or 'user'
    return f'{prefix}{base}'[:140]


def generated_password() -> str:
    alphabet = string.ascii_letters + string.digits + '!@#$%^'
    while True:
        candidate = ''.join(secrets.choice(alphabet) for _ in range(18))
        if (
            any(ch.islower() for ch in candidate)
            and any(ch.isupper() for ch in candidate)
            and any(ch.isdigit() for ch in candidate)
        ):
            return candidate


def required_row(sheet_name: str, row: dict) -> bool:
    if sheet_name == 'SOTIB OLISH':
        return (
            all(not is_blank(pick(row, key)) for key in ('SANA', 'MAHSULOT', 'SONI', 'MAHSULOT NARXI'))
            and dec(pick(row, 'SONI')) > ZERO
            and dec(pick(row, 'MAHSULOT NARXI')) > ZERO
        )
    if sheet_name == 'STOCK TRANSFER':
        return all(not is_blank(pick(row, key)) for key in ('SANA', 'MAHSULOT', 'SONI', 'KIRIM', 'CHIQIM'))
    if sheet_name == 'SOTUV':
        return (
            all(not is_blank(pick(row, key)) for key in ('SOTUV SANASI', 'MAHSULOT', 'JAMI DONA', 'SOTUV NARXI'))
            and dec(pick(row, 'JAMI DONA')) > ZERO
            and dec(pick(row, 'SOTUV NARXI')) > ZERO
        )
    if sheet_name == 'TUSHUM':
        return all(not is_blank(pick(row, key)) for key in ('SANA', 'MIJOZ', 'MIQDOR'))
    if sheet_name == 'XARAJAT':
        return all(not is_blank(pick(row, key)) for key in ('SANA', 'MIQDOR'))
    if sheet_name == 'PUL AYRIBOSHLASH':
        return all(not is_blank(pick(row, key)) for key in ('SANA', 'KIRIM', 'CHIQIM', 'CHIQIM MIQDORI'))
    return any(not is_blank(value) for value in row.values())


def sheet_rows(workbook, sheet_name: str) -> list[dict]:
    if sheet_name not in workbook.sheetnames:
        return []
    ws = workbook[sheet_name]
    headers = [
        str(cell.value).strip() if cell.value is not None else ''
        for cell in next(ws.iter_rows(min_row=1, max_row=1))
    ]
    if sheet_name == 'TUSHUM' and (not headers[0] or norm(headers[0]) not in {'SANA', 'DATE'}):
        headers[0] = 'SANA'

    result = []
    for row_id, cells in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        row = {header: cell.value for header, cell in zip(headers, cells) if header}
        if all(is_blank(value) for value in row.values()):
            continue
        if not required_row(sheet_name, row):
            continue
        row['_row_id'] = row_id
        result.append(row)
    return result


def find_header(ws, first_label: str) -> int:
    target = norm(first_label)
    for row in range(1, min(ws.max_row or 1, 80) + 1):
        if norm(ws.cell(row, 1).value) == target:
            return row
    raise CommandError(f'Header {first_label!r} not found in {ws.title}.')


def master_sheet_rows(workbook, sheet_name: str, first_label: str) -> list[dict]:
    ws = workbook[sheet_name]
    header_row = find_header(ws, first_label)
    headers = [ws.cell(header_row, col).value for col in range(1, (ws.max_column or 0) + 1)]
    rows = []
    for row_id in range(header_row + 1, (ws.max_row or 1) + 1):
        row = {
            headers[index]: ws.cell(row_id, index + 1).value
            for index in range(len(headers))
            if headers[index]
        }
        if any(not is_blank(value) for value in row.values()):
            row['_row_id'] = row_id
            rows.append(row)
    return rows


def read_foyda_terms(workbook) -> tuple[str, list[AgreementTerm]]:
    if 'FOYDA_TAQSIMOTI' not in workbook.sheetnames:
        raise CommandError('FOYDA_TAQSIMOTI sheet is required in every local deal file.')
    ws = workbook['FOYDA_TAQSIMOTI']
    capital_header = str(ws.cell(1, 2).value or '')
    terms_currency = 'UZS' if 'UZS' in capital_header.upper() else 'USD'
    terms: list[AgreementTerm] = []
    for row in range(2, min(ws.max_row or 2, 12) + 1):
        name = display_name(ws.cell(row, 1).value)
        capital = ws.cell(row, 2).value
        profit_share = ws.cell(row, 4).value
        if not name or capital in (None, '') or profit_share in (None, ''):
            continue
        if dec(profit_share, Decimal('-1')) < ZERO:
            continue
        terms.append(AgreementTerm(
            name=name,
            capital=money(capital),
            capital_share=ratio(ws.cell(row, 3).value),
            profit_share=ratio(profit_share),
        ))
    return terms_currency, terms


def account_amount(row: dict) -> tuple[str, str, Decimal]:
    account = norm(pick(row, "TO'LOV TURI"))
    row_currency = currency(pick(row, 'VALYUTA'))
    if account == 'KASSA DOLLAR':
        amount = pick(row, 'JAMI (USD)', 'JAMI\n(USD)')
        return account, 'USD', money(amount if not is_blank(amount) else pick(row, 'MIQDOR'))
    if account in {'KASSA SOM', 'PLASTIK SOM'}:
        amount = pick(row, 'JAMI (SOM)', 'JAMI\n(SOM)')
        return account, 'UZS', money(amount if not is_blank(amount) else pick(row, 'MIQDOR'))
    return account, row_currency, money(pick(row, 'MIQDOR'))


class Command(BaseCommand):
    help = 'Preflight and safely apply Sherik Excel Replay sources.'

    def add_arguments(self, parser):
        parser.add_argument('--folder', default=str(DEFAULT_FOLDER))
        parser.add_argument('--dry-run', action='store_true')
        parser.add_argument('--apply', action='store_true')
        parser.add_argument('--wipe', action='store_true')
        parser.add_argument('--stage', choices=['snapshot', 'plan', 'preflight'], default='preflight')
        parser.add_argument(
            '--backup-dir',
            default=str(Path.home() / '.cache' / 'microposs' / 'sherik_replay' / 'backups'),
        )
        parser.add_argument('--admin-username', default='sherik_admin')
        parser.add_argument('--admin-password', default='')
        parser.add_argument('--default-user-password', default='')
        parser.add_argument(
            '--confirm-production-apply',
            action='store_true',
            help='Allow replay apply when DEBUG=False.',
        )

    def handle(self, *args, **options):
        if bool(options['dry_run']) == bool(options['apply']):
            raise CommandError('Choose exactly one: --dry-run or --apply.')
        if options['apply']:
            require_debug_or_confirmation(
                command_name='sherik_excel_replay',
                confirmed=options['confirm_production_apply'],
                flag_name='--confirm-production-apply',
                action='this command can wipe and replay Sherik Excel data',
            )
            if not options['wipe']:
                raise CommandError('--apply requires --wipe so the replay starts from a clean database.')

        folder = Path(options['folder']).expanduser()
        if not folder.exists():
            raise CommandError(f'Folder not found: {folder}')
        master_path = folder / MASTER_FILE
        if not master_path.exists():
            raise CommandError(f'Master workbook not found: {master_path}')

        master = self._read_master(master_path)
        local_deals = self._read_local_deals(folder)
        report = self._build_report(master, local_deals)
        self._print_report(report, stage=options['stage'])

        if report['blockers']:
            raise CommandError(
                f"Sherik Excel Replay preflight has {len(report['blockers'])} blocker(s); "
                'database was not changed.'
            )
        if options['dry_run']:
            return

        backup_path = self._backup_database(Path(options['backup_dir']).expanduser())
        self.stdout.write(f'Backup written: {backup_path}')

        admin_password = options['admin_password'] or generated_password()
        default_password = options['default_user_password'] or generated_password()
        with transaction.atomic():
            self._wipe_database()
            summary = self._apply_replay(
                master=master,
                local_deals=local_deals,
                report=report,
                admin_username=options['admin_username'],
                admin_password=admin_password,
                default_user_password=default_password,
            )

        credentials_path = self._write_credentials(
            backup_path.parent,
            admin_username=options['admin_username'],
            admin_password=admin_password,
            default_user_password=default_password,
            summary=summary,
        )
        self.stdout.write(self.style.SUCCESS('Sherik Excel Replay apply complete.'))
        self.stdout.write(f'Credentials written: {credentials_path}')

    def _read_master(self, path: Path) -> dict:
        workbook = load_workbook(path, data_only=True, read_only=False)
        investors = []
        if 'INVESTORLAR' in workbook.sheetnames:
            for row in master_sheet_rows(workbook, 'INVESTORLAR', 'Investor ID'):
                name = display_name(row.get('Investor nomi'))
                if name:
                    investors.append({
                        'id': display_name(row.get('Investor ID')),
                        'name': name,
                        'phone': display_name(row.get('Telefon')),
                        'status': display_name(row.get('Status')),
                        'row_id': row['_row_id'],
                    })

        deals = {
            str(row.get('Bitim ID') or '').strip(): row
            for row in master_sheet_rows(workbook, 'BITIMLAR', 'Bitim ID')
            if str(row.get('Bitim ID') or '').strip()
        }

        funding_by_deal = defaultdict(list)
        source_corrections = []
        for row in master_sheet_rows(workbook, 'BITIM_FUNDING', 'Sana'):
            deal_id = str(row.get('Bitim ID') or '').strip()
            investor = display_name(row.get('Investor'))
            amount = money(row.get('Tikilgan summa USD'))
            if deal_id and investor and amount:
                effective_date = parse_dt(row.get('Sana'))
                source_date = effective_date
                for spec in SOURCE_CORRECTIONS.get(deal_id, {}).get('funding_date_overrides', []):
                    if int(row['_row_id']) != int(spec['row_id']):
                        continue
                    expected_investor = clean_person_name(spec['investor'])
                    if clean_person_name(investor) != expected_investor:
                        raise CommandError(
                            f"{deal_id}: funding correction row {spec['row_id']} expected "
                            f"{expected_investor}, got {investor}."
                        )
                    effective_date = parse_dt(spec['effective_date'])
                    source_corrections.append(
                        f"BITIM_FUNDING row {spec['row_id']}: effective date "
                        f"{date_label(source_date)} -> {date_label(effective_date)} "
                        f"for {investor} ({spec['reason']})"
                    )
                    break
                funding_by_deal[deal_id].append({
                    'date': effective_date,
                    'source_date': source_date,
                    'investor': investor,
                    'amount_usd': amount,
                    'funding_source': display_name(row.get('Funding source')),
                    'row_id': row['_row_id'],
                })

        payouts_by_deal = defaultdict(list)
        if 'INVESTOR_PAYOUTS' in workbook.sheetnames:
            for row in master_sheet_rows(workbook, 'INVESTOR_PAYOUTS', 'Sana'):
                deal_id = str(row.get('Bitim ID') or '').strip()
                amount = money(row.get('USD summa'))
                if deal_id and amount:
                    payouts_by_deal[deal_id].append({
                        'date': row.get('Sana'),
                        'investor': display_name(row.get('Investor')),
                        'kind': display_name(row.get('Payout turi')),
                        'currency': currency(row.get('Valyuta')),
                        'amount': amount,
                        'row_id': row['_row_id'],
                    })

        withdrawals_by_deal = defaultdict(list)
        if 'BITIMDAN_CHIQARILGAN_PUL' in workbook.sheetnames:
            for row in master_sheet_rows(workbook, 'BITIMDAN_CHIQARILGAN_PUL', 'Sana'):
                deal_id = str(row.get('Bitim ID') or '').strip()
                amount = money(row.get('Chiqarilgan summa USD'))
                if deal_id and amount:
                    withdrawals_by_deal[deal_id].append({
                        'date': row.get('Sana'),
                        'amount': amount,
                        'source': display_name(row.get('Manba')),
                        'row_id': row['_row_id'],
                    })

        return {
            'investors': investors,
            'deals': deals,
            'funding_by_deal': dict(funding_by_deal),
            'payouts_by_deal': dict(payouts_by_deal),
            'withdrawals_by_deal': dict(withdrawals_by_deal),
            'source_corrections': source_corrections,
        }

    def _read_local_deals(self, folder: Path) -> dict[str, LocalDeal]:
        local_deals = {}
        for path in sorted(folder.glob('*.xlsx')):
            if path.name == MASTER_FILE:
                continue
            deal_id = path.stem.strip()
            workbook = load_workbook(path, data_only=True, read_only=False)
            terms_currency, terms = read_foyda_terms(workbook)
            rows = {
                sheet: sheet_rows(workbook, sheet)
                for sheet in (
                    'SOTIB OLISH',
                    'STOCK TRANSFER',
                    'SOTUV',
                    'TUSHUM',
                    'XARAJAT',
                    'PUL AYRIBOSHLASH',
                )
            }
            source_corrections = apply_source_corrections(deal_id, rows)
            local_deals[deal_id] = LocalDeal(
                deal_id=deal_id,
                path=path,
                terms_currency=terms_currency,
                terms=terms,
                rows=rows,
                source_corrections=source_corrections,
            )
        return local_deals

    def _backup_database(self, backup_dir: Path) -> Path:
        backup_dir.mkdir(parents=True, exist_ok=True)
        path = backup_dir / f'sherik-replay-before-{timezone.now().strftime("%Y%m%d-%H%M%S")}.json'
        with path.open('w', encoding='utf-8') as handle:
            call_command('dumpdata', stdout=handle, indent=2)
        return path

    def _wipe_database(self) -> None:
        cache.clear()
        keep = {'auth', 'admin', 'contenttypes', 'sessions', 'django_celery_beat'}
        tables = [
            model._meta.db_table
            for model in django_apps.get_models()
            if model._meta.app_label not in keep
        ]
        if tables:
            quoted = ', '.join(f'"{table}"' for table in tables)
            with connection.cursor() as cursor:
                cursor.execute(f'TRUNCATE {quoted} RESTART IDENTITY CASCADE;')
        User.objects.all().delete()

    def _write_credentials(
        self,
        folder: Path,
        *,
        admin_username: str,
        admin_password: str,
        default_user_password: str,
        summary: dict,
    ) -> Path:
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f'sherik-replay-credentials-{timezone.now().strftime("%Y%m%d-%H%M%S")}.txt'
        lines = [
            'Sherik Excel Replay credentials',
            '',
            f'admin_username={admin_username}',
            f'admin_password={admin_password}',
            f'default_user_password={default_user_password}',
            '',
            'business_users:',
        ]
        for username in sorted(summary.get('business_users', [])):
            lines.append(f'  {username}')
        lines.append('investor_users:')
        for username in sorted(summary.get('investor_users', [])):
            lines.append(f'  {username}')
        path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
        return path

    def _apply_replay(
        self,
        *,
        master: dict,
        local_deals: dict[str, LocalDeal],
        report: dict,
        admin_username: str,
        admin_password: str,
        default_user_password: str,
    ) -> dict:
        admin = self._create_admin(admin_username, admin_password)
        investor_users = self._create_investor_users(master, default_user_password)
        contexts = self._create_business_contexts(master, local_deals, admin, default_user_password)
        self._seed_workbook_fx_rates(contexts, local_deals, report)
        self._connect_investors(contexts, investor_users, report)
        self._create_catalog(contexts, local_deals, report)
        replay_summary = self._apply_timeline(contexts, master, local_deals, report)
        return {
            'business_users': [ctx['owner'].username for ctx in contexts.values()],
            'investor_users': [user.username for user in investor_users.values()],
            **replay_summary,
        }

    def _seed_workbook_fx_rates(
        self,
        contexts: dict[str, dict],
        local_deals: dict[str, LocalDeal],
        report: dict,
    ) -> None:
        from apps.finance.fx_rates import upsert_exchange_rate
        from apps.finance.models import ExchangeRate

        agreements = {row['deal_id']: row for row in report['agreements']}
        rates_by_business: dict[str, dict] = defaultdict(lambda: defaultdict(Counter))
        for deal_id, deal in local_deals.items():
            agreement = agreements.get(deal_id, {})
            business = business_key(deal_id, clean_person_name(agreement.get('operator_name') or ''))
            for rate_date, rate in self._workbook_fx_rate_rows(deal):
                rates_by_business[business][rate_date][rate] += 1

        for business, date_map in rates_by_business.items():
            ctx = contexts.get(business)
            if ctx is None:
                raise CommandError(f'{business}: cannot seed workbook FX rates without business context.')
            for rate_date, counter in sorted(date_map.items()):
                rate = counter.most_common(1)[0][0]
                upsert_exchange_rate(
                    tenant_id=ctx['tenant_id'],
                    base_currency='USD',
                    quote_currency='UZS',
                    rate_date=rate_date,
                    rate=rate,
                    source=ExchangeRate.Source.MANUAL,
                    is_manual=True,
                    notes='Sherik replay workbook FX fallback',
                    overwrite_manual=False,
                )

    def _workbook_fx_rate_rows(self, deal: LocalDeal):
        def add_from_rows(rows: list[dict], date_reader):
            for row in rows:
                rate = dec(pick(row, 'KURS'), ZERO)
                if rate <= Decimal('1'):
                    continue
                yield date_reader(row).date(), rate

        yield from add_from_rows(deal.rows['SOTIB OLISH'], self._purchase_received_at)
        yield from add_from_rows(deal.rows['SOTUV'], lambda row: parse_dt(pick(row, 'SOTUV SANASI')))
        yield from add_from_rows(deal.rows['TUSHUM'], lambda row: parse_dt(pick(row, 'SANA')))
        yield from add_from_rows(deal.rows['XARAJAT'], lambda row: parse_dt(pick(row, 'SANA')))
        yield from add_from_rows(deal.rows['PUL AYRIBOSHLASH'], lambda row: parse_dt(pick(row, 'SANA')))

    def _create_admin(self, username: str, password: str) -> User:
        groups = {
            name: Group.objects.get_or_create(name=name)[0]
            for name in ('owner', 'cashier', 'warehouse', 'investor')
        }
        admin = User.objects.create_user(
            username=username,
            email=f'{username}@local.dev',
            password=password,
        )
        admin.is_active = True
        admin.is_staff = True
        admin.is_superuser = True
        admin.save(update_fields=['is_active', 'is_staff', 'is_superuser'])
        admin.groups.set([groups['owner']])
        return admin

    def _create_investor_users(self, master: dict, password: str) -> dict[str, User]:
        Group.objects.get_or_create(name='investor')
        investor_group = Group.objects.get(name='investor')
        users = {}
        used = set(User.objects.values_list('username', flat=True))
        for investor in master.get('investors', []):
            name = clean_person_name(investor['name'])
            username = self._unique_username(slug(name, prefix='inv_'), used)
            user = User.objects.create_user(
                username=username,
                email=f'{username}@local.dev',
                password=password,
            )
            user.is_active = True
            user.save(update_fields=['is_active'])
            user.groups.set([investor_group])
            from apps.core.services import get_or_create_investment_profile

            get_or_create_investment_profile(user, display_name=name)
            users[name] = user
        return users

    def _create_business_contexts(
        self,
        master: dict,
        local_deals: dict[str, LocalDeal],
        admin: User,
        password: str,
    ) -> dict[str, dict]:
        from apps.core.services import (
            approve_business_registration_request,
            create_business_registration_request,
        )

        used = set(User.objects.values_list('username', flat=True))
        contexts = {}
        for deal_id, deal in sorted(local_deals.items()):
            agreement = self._agreement_report(deal_id, deal, master, [], [])
            business = business_key(deal_id, clean_person_name(agreement.get('operator_name') or ''))
            if business in contexts:
                continue
            username = self._unique_username(slug(business, prefix='biz_'), used)
            request = create_business_registration_request(
                username=username,
                password=password,
                first_name=business,
                last_name='',
                phone=f'+998000{len(contexts) + 1:06d}',
                business_name=business,
            )
            request = approve_business_registration_request(
                request_id=request.id,
                reviewed_by=admin,
            )
            business_obj = request.approved_business
            owner = request.approved_user
            contexts[business] = self._create_operational_baseline(
                business=business_obj,
                owner=owner,
                display_name=business,
            )
        return contexts

    def _create_operational_baseline(self, *, business, owner: User, display_name: str) -> dict:
        from apps.catalog.models import Category, DiscountReason
        from apps.core.demo_constants import DEFAULT_DEMO_USD_UZS_RATE
        from apps.core.models import Partner
        from apps.finance.chart_of_accounts import setup_chart_of_accounts
        from apps.finance.models import Account, CashAccount, ExchangeRate
        from apps.inventory.models import Warehouse
        from apps.suppliers.models import Supplier

        setup_chart_of_accounts(business.id)
        ExchangeRate.objects.update_or_create(
            tenant=business,
            base_currency='USD',
            quote_currency='UZS',
            rate_date=timezone.localdate(),
            defaults={
                'rate': DEFAULT_DEMO_USD_UZS_RATE,
                'source': ExchangeRate.Source.MANUAL,
                'is_manual': True,
                'notes': 'Sherik replay fallback FX rate',
                'raw_payload': {},
                'fetched_at': timezone.now(),
            },
        )
        cash_account = Account.objects.get(tenant_id=business.id, code='1000')
        bank_account = Account.objects.get(tenant_id=business.id, code='1010')
        cash_accounts = {}
        for name, cur, kind, linked in (
            ('KASSA SOM', 'UZS', CashAccount.Kind.CASH, cash_account),
            ('KASSA DOLLAR', 'USD', CashAccount.Kind.CASH, bank_account),
            ('PLASTIK SOM', 'UZS', CashAccount.Kind.CARD_TERMINAL, bank_account),
        ):
            cash_accounts[name] = CashAccount.objects.create(
                tenant=business,
                name=name,
                currency=cur,
                kind=kind,
                balance=ZERO,
                is_active=True,
                linked_account=linked,
            )
        warehouses = {
            'ASOSIY': Warehouse.objects.create(
                tenant=business,
                name='Основной склад',
                kind=Warehouse.WarehouseKind.STORAGE,
                is_active=True,
            ),
            'DOKON': Warehouse.objects.create(
                tenant=business,
                name='Основной магазин',
                kind=Warehouse.WarehouseKind.SHOP,
                is_active=True,
            ),
        }
        operator = Partner.objects.get(tenant=business, role=Partner.Role.OPERATOR)
        if operator.display_name != display_name:
            operator.display_name = display_name
            operator.save(update_fields=['display_name', 'updated_at'])
        category = Category.objects.create(tenant=business, name='Excel товары', sort_order=1)
        discount_reason = DiscountReason.objects.create(
            tenant=business,
            name='Торг',
            is_default=True,
            is_active=True,
        )
        supplier = Supplier.objects.create(tenant=business, name='Excel supplier', is_active=True)
        return {
            'business': business,
            'tenant_id': business.id,
            'owner': owner,
            'operator': operator,
            'cash_accounts': cash_accounts,
            'warehouses': warehouses,
            'category': category,
            'discount_reason': discount_reason,
            'supplier': supplier,
            'investor_partners': {},
            'variants': {},
        }

    def _connect_investors(self, contexts: dict[str, dict], investor_users: dict[str, User], report: dict) -> None:
        from apps.core.services import accept_investor_invite, create_investor_invite
        from apps.investors.models import Investor

        investors_by_business = defaultdict(set)
        for agreement in report['agreements']:
            business = business_key(agreement['deal_id'], clean_person_name(agreement.get('operator_name') or ''))
            for split in agreement.get('funding_split') or []:
                investors_by_business[business].add(clean_person_name(split['investor']))

        for business, investors in sorted(investors_by_business.items()):
            ctx = contexts[business]
            for investor in sorted(investors):
                user = investor_users.get(investor)
                if user is None:
                    raise CommandError(f'{business}: investor is in funding but missing from INVESTORLAR: {investor}')
                invite = create_investor_invite(
                    tenant_id=ctx['tenant_id'],
                    invited_by_id=ctx['owner'].id,
                    email=user.email,
                    display_name=investor,
                    expires_days=60,
                )
                _invite, relation = accept_investor_invite(
                    token=str(invite.token),
                    user=user,
                    display_name=investor,
                )
                partner = relation.partner
                ctx['investor_partners'][investor] = partner
                Investor.objects.update_or_create(
                    tenant=ctx['business'],
                    user=user,
                    defaults={
                        'name': investor,
                        'email': user.email,
                        'is_active': True,
                    },
                )

    def _create_catalog(self, contexts: dict[str, dict], local_deals: dict[str, LocalDeal], report: dict) -> None:
        from apps.catalog.services import create_product_with_variants

        purchase_products = defaultdict(set)
        sale_price_modes = defaultdict(lambda: defaultdict(Counter))
        for agreement in report['agreements']:
            deal_id = agreement['deal_id']
            deal = local_deals[deal_id]
            business = business_key(deal_id, clean_person_name(agreement.get('operator_name') or ''))
            for row in deal.rows['SOTIB OLISH']:
                key = product_key(pick(row, 'MAHSULOT'))
                if key:
                    purchase_products[business].add(key)
            for row in deal.rows['SOTUV']:
                key = product_key(pick(row, 'MAHSULOT'))
                if not key:
                    continue
                row_currency = currency(pick(row, 'VALYUTA'))
                price = dec(pick(row, 'SOTUV NARXI'))
                fx = dec(pick(row, 'KURS'), Decimal('1'))
                price_uzs = money(price if row_currency == 'UZS' else price * fx)
                sale_price_modes[business][key][price_uzs] += 1

        for business, products in sorted(purchase_products.items()):
            ctx = contexts[business]
            for key in sorted(products):
                mode = sale_price_modes[business][key].most_common(1)
                base_price = mode[0][0] if mode else Decimal('100000.00')
                product = create_product_with_variants(
                    tenant_id=ctx['tenant_id'],
                    name=key,
                    category_id=ctx['category'].id,
                    base_price=str(base_price),
                    pricing_mode='EDITABLE',
                    variant_data=None,
                )
                ctx['variants'][key] = product.variants.filter(is_active=True).first()

    def _apply_timeline(
        self,
        contexts: dict[str, dict],
        master: dict,
        local_deals: dict[str, LocalDeal],
        report: dict,
    ) -> dict:
        agreements = {row['deal_id']: row for row in report['agreements']}
        state = {'deals': {}, 'sessions': {}}
        events = []
        for deal_id, deal in sorted(local_deals.items()):
            agreement = agreements[deal_id]
            business = business_key(deal_id, clean_person_name(agreement.get('operator_name') or ''))
            setup_at = self._deal_start_at(master, deal_id, deal)
            events.append((business, setup_at, 10, deal_id, 0, 'setup_deal', None))
            for row in agreement.get('supplemental_funding') or []:
                events.append((
                    business,
                    parse_dt(row['date']),
                    15,
                    deal_id,
                    int(row.get('row_id') or 0),
                    'supplemental_contribution',
                    row,
                ))
            for row in deal.rows['SOTIB OLISH']:
                events.append((
                    business,
                    self._purchase_received_at(row),
                    20,
                    deal_id,
                    int(row.get('_row_id') or 0),
                    'receive',
                    row,
                ))
            for row in deal.rows['STOCK TRANSFER']:
                qty = dec(pick(row, 'SONI'))
                if qty > ZERO:
                    events.append((
                        business,
                        parse_dt(pick(row, 'SANA')),
                        30,
                        deal_id,
                        int(row.get('_row_id') or 0),
                        'transfer',
                        row,
                    ))
            for row in deal.rows['SOTUV']:
                events.append((
                    business,
                    parse_dt(pick(row, 'SOTUV SANASI')),
                    40,
                    deal_id,
                    int(row.get('_row_id') or 0),
                    'sale',
                    row,
                ))

        for business, _dt, _order, deal_id, _row_id, kind, row in sorted(events, key=lambda item: item[:5]):
            ctx = contexts[business]
            if kind == 'setup_deal':
                self._setup_deal(ctx, master, local_deals[deal_id], agreements[deal_id], state)
            elif kind == 'supplemental_contribution':
                self._apply_supplemental_contribution(
                    ctx,
                    local_deals[deal_id],
                    agreements[deal_id],
                    state,
                    row,
                )
            elif kind == 'receive':
                self._receive_purchase_row(ctx, local_deals[deal_id], agreements[deal_id], state, row)
            elif kind == 'transfer':
                self._transfer_row(ctx, state, row)
            elif kind == 'sale':
                self._sale_row(ctx, state, row)

        self._close_sessions(state)
        return {
            'agreements_applied': len(state['deals']),
            'open_sessions_closed': len(state['sessions']),
            'capital_rollovers': sum(
                len(row.get('capital_rollovers', []))
                for row in state['deals'].values()
            ),
        }

    def _setup_deal(self, ctx: dict, master: dict, deal: LocalDeal, agreement_report: dict, state: dict) -> None:
        from apps.partnerships.workspace import create_workspace, dispatch_workspace_action
        from apps.partnerships.workspace_funding import convert_workspace_capital_pool, link_workspace_agreement
        from apps.partnerships.workspace_support import add_agreement_contribution, create_investment_agreement

        purchase_rows = deal.rows['SOTIB OLISH']
        if not purchase_rows:
            raise CommandError(f'{deal.deal_id}: no purchase rows.')

        purchase_currencies = {currency(pick(row, 'PUL BIRLIGI')) for row in purchase_rows}
        if len(purchase_currencies) != 1:
            raise CommandError(f'{deal.deal_id}: mixed purchase currencies are not replay-ready.')
        purchase_currency = next(iter(purchase_currencies))
        purchase_total = self._purchase_total(purchase_rows)
        start_at = self._deal_start_at(master, deal.deal_id, deal)
        agreement_currency = 'USD' if agreement_report['master_pool'] > ZERO else deal.terms_currency

        partners, contribution_rows = self._agreement_partners_payload(ctx, agreement_report, agreement_currency)
        planned_budget = money(sum((money(row['amount']) for row in contribution_rows), ZERO))
        if planned_budget <= ZERO:
            planned_budget = money(agreement_report['total_capital'])

        procurement = create_workspace(
            tenant_id=ctx['tenant_id'],
            funding_source='PARTNERSHIP',
            primary_currency=purchase_currency,
            supplier_id=ctx['supplier'].id,
            notes=f'Sherik Excel Replay {deal.deal_id}',
            client_request_id=None,
        )
        agreement = create_investment_agreement(
            tenant_id=ctx['tenant_id'],
            opened_at=start_at,
            supplier_id=ctx['supplier'].id,
            mudaraba_ratio=agreement_report['mudaraba_ratio'],
            planned_budget=planned_budget,
            currency=agreement_currency,
            notes=f'Sherik Excel Replay agreement {deal.deal_id}',
            created_by_id=ctx['owner'].id,
            actor_partner_id=ctx['operator'].id,
            reconciliation_mode='AGREED',
            review_at=start_at,
            offline_agreed_at=start_at,
            offline_agreement_reference=deal.deal_id,
            payout_policy={
                'review_interval_days': 30,
                'minimum_available_amount': Decimal('1.00'),
                'minimum_days_between_payouts': 0,
                'reserve_amount': Decimal('0.00'),
                'allow_partial': True,
                'trigger_mode': 'ANY',
            },
            partners=partners,
        )
        link_workspace_agreement(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            agreement_id=agreement.id,
        )
        procurement.refresh_from_db()

        for row in contribution_rows:
            amount = money(row['amount'])
            if amount <= ZERO:
                continue
            add_agreement_contribution(
                tenant_id=ctx['tenant_id'],
                agreement_id=agreement.id,
                partner_id=row['partner_id'],
                amount=amount,
                currency=agreement_currency,
                fx_rate=Decimal('1'),
                date=start_at,
                notes=f'Sherik replay capital {deal.deal_id}',
                created_by_id=ctx['owner'].id,
                actor_partner_id=row['partner_id'],
            )

        if purchase_currency != agreement_currency:
            implied_fx = agreement_report.get('implied_fx')
            if not implied_fx:
                raise CommandError(f'{deal.deal_id}: purchase currency differs from agreement currency but no implied FX was resolved.')
            initial_conversion_amount = money(sum(
                (money(row['amount']) for row in contribution_rows),
                ZERO,
            ))
            if initial_conversion_amount > ZERO:
                convert_workspace_capital_pool(
                    tenant_id=ctx['tenant_id'],
                    procurement=procurement,
                    payload={
                        'to_currency': purchase_currency,
                        'from_amount': str(initial_conversion_amount),
                        'rate': str(implied_fx),
                        'date': start_at,
                    },
                )

        items_payload = []
        for row in purchase_rows:
            key = product_key(pick(row, 'MAHSULOT'))
            variant = ctx['variants'].get(key)
            if variant is None:
                raise CommandError(f'{deal.deal_id}: product was not created: {key}')
            fx = dec(pick(row, 'KURS'), Decimal('1')) if purchase_currency == 'USD' else Decimal('1')
            items_payload.append({
                'product_variant_id': variant.id,
                'quantity': str(dec(pick(row, 'SONI'))),
                'unit_purchase_price': str(dec(pick(row, 'MAHSULOT NARXI'))),
                'currency': purchase_currency,
                'fx_rate': str(fx),
            })
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='UPDATE_ITEMS',
            payload={'payload': {'items': items_payload}},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        item_by_row_id = {
            int(row.get('_row_id') or 0): item
            for row, item in zip(purchase_rows, procurement.items.order_by('id'))
        }
        settlement_fx = self._purchase_fx(purchase_rows[0], purchase_currency)
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=procurement,
            action='UPDATE_SETTLEMENT',
            payload={'payload': {
                'type': 'AT_RECEIPT',
                'currency_of_obligation': purchase_currency,
                'fx_rate_at_obligation': str(settlement_fx),
                'total_amount_due': str(purchase_total),
                'notes': f'Sherik replay settlement {deal.deal_id}',
            }},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        state['deals'][deal.deal_id] = {
            'procurement': procurement,
            'agreement': agreement,
            'item_by_row_id': item_by_row_id,
            'agreement_currency': agreement_currency,
            'purchase_currency': purchase_currency,
            'implied_fx': agreement_report.get('implied_fx'),
            'purchase_fx': settlement_fx,
            'partners': partners,
            'capital_remaining': {
                int(row['partner_id']): money(row['planned_capital_share'])
                for row in partners
            },
            'capital_rollovers': [],
        }
        self._sync_replay_capital_remaining(state['deals'][deal.deal_id])

    def _apply_supplemental_contribution(
        self,
        ctx: dict,
        deal: LocalDeal,
        agreement_report: dict,
        state: dict,
        row: dict,
    ) -> None:
        from apps.partnerships.workspace_funding import convert_workspace_capital_pool
        from apps.partnerships.workspace_support import add_agreement_contribution

        deal_state = state['deals'].get(deal.deal_id)
        if deal_state is None:
            raise CommandError(f'{deal.deal_id}: supplemental funding cannot run before agreement setup.')
        investor = clean_person_name(row['investor'])
        partner = ctx['investor_partners'].get(investor)
        if partner is None:
            raise CommandError(f'{deal.deal_id}: supplemental funding investor is not linked: {investor}.')
        date = parse_dt(row['date'])
        agreement_currency = deal_state['agreement_currency']
        amount = self._funding_amount_for_currency(row, agreement_currency)
        if amount <= ZERO:
            return

        add_agreement_contribution(
            tenant_id=ctx['tenant_id'],
            agreement_id=deal_state['agreement'].id,
            partner_id=partner.id,
            amount=amount,
            currency=agreement_currency,
            fx_rate=Decimal('1'),
            date=date,
            notes=(
                f"Sherik replay supplemental funding {deal.deal_id} "
                f"row {row.get('row_id')} source={row.get('funding_source') or 'Master'}"
            ),
            client_request_id=str(uuid5(
                NAMESPACE_URL,
                f"sherik-supplemental-funding:{deal.deal_id}:{row.get('row_id')}:{partner.id}:{amount}",
            )),
            created_by_id=ctx['owner'].id,
            actor_partner_id=partner.id,
        )

        if deal_state['purchase_currency'] != agreement_currency:
            implied_fx = agreement_report.get('implied_fx')
            if not implied_fx:
                raise CommandError(
                    f'{deal.deal_id}: supplemental funding needs conversion but no implied FX was resolved.'
                )
            convert_workspace_capital_pool(
                tenant_id=ctx['tenant_id'],
                procurement=deal_state['procurement'],
                payload={
                    'to_currency': deal_state['purchase_currency'],
                    'from_amount': str(amount),
                    'rate': str(implied_fx),
                    'date': date,
                },
            )
        self._sync_replay_capital_remaining(deal_state)

    def _receive_purchase_row(self, ctx: dict, deal: LocalDeal, agreement_report: dict, state: dict, row: dict) -> None:
        from apps.partnerships.workspace import dispatch_workspace_action

        deal_state = state['deals'][deal.deal_id]
        item = deal_state['item_by_row_id'].get(int(row.get('_row_id') or 0))
        if item is None:
            raise CommandError(f"{deal.deal_id}: purchase item not found for row {row.get('_row_id')}.")
        received_at = self._purchase_received_at(row)
        warehouse = ctx['warehouses'][self._purchase_warehouse_key(row)]
        required_base = self._row_required_base(row, deal_state)
        source_ref = f"purchase-row:{int(row.get('_row_id') or 0)}"
        self._ensure_replay_capital(ctx, deal, deal_state, required_base, received_at, source_ref=source_ref)
        self._sync_replay_capital_remaining(deal_state)
        try:
            allocations = self._capital_allocations_for_amount(
                deal_state['partners'],
                required_base,
                deal_state['capital_remaining'],
            )
        except CommandError as exc:
            remaining = ', '.join(
                f"{row['partner_id']}={fmt(deal_state['capital_remaining'].get(int(row['partner_id']), ZERO))}"
                for row in deal_state['partners']
            )
            raise CommandError(
                f"{deal.deal_id}: cannot allocate purchase row {row.get('_row_id')} "
                f"{product_key(pick(row, 'MAHSULOT'))} at {received_at.isoformat()} "
                f"for {fmt(required_base)} {deal_state['agreement_currency']}. "
                f"Remaining by partner: {remaining}. Root: {exc}"
            ) from exc
        procurement = dispatch_workspace_action(
            tenant_id=ctx['tenant_id'],
            procurement=deal_state['procurement'],
            action='RECEIVE_BATCH',
            payload={'payload': {
                'warehouse_id': warehouse.id,
                'item_ids': [item.id],
                'capital_allocations': allocations,
                'received_at': received_at.isoformat(),
            }},
            user_id=ctx['owner'].id,
        )
        procurement.refresh_from_db()
        deal_state['procurement'] = procurement

    def _ensure_replay_capital(
        self,
        ctx: dict,
        deal: LocalDeal,
        deal_state: dict,
        required: Decimal,
        date,
        *,
        source_ref: str,
    ) -> None:
        self._sync_replay_capital_remaining(deal_state)
        total_remaining = money(sum(deal_state['capital_remaining'].values(), ZERO))
        required = money(required)
        if required - total_remaining <= CENT:
            return

        from apps.core.demo_constants import DEFAULT_DEMO_USD_UZS_RATE
        from apps.finance.fx_rates import upsert_exchange_rate
        from apps.finance.models import ExchangeRate
        from apps.partnerships.lifecycle_services import build_payout_decision_preview, execute_payout_decision
        from apps.partnerships.models import PayoutDecision
        from apps.partnerships.workspace_support import add_agreement_contribution

        shortage = money(required - total_remaining)
        topups = self._capital_topups_for_amount(deal_state['partners'], shortage)
        role_by_partner = {
            int(row['partner_id']): str(row.get('role') or '').upper()
            for row in deal_state['partners']
        }
        investor_topup = money(sum(
            money(row['amount'])
            for row in topups
            if role_by_partner.get(int(row['partner_id'])) == 'INVESTOR'
        ))
        operator_topups = [
            row for row in topups
            if role_by_partner.get(int(row['partner_id'])) != 'INVESTOR'
            and money(row['amount']) > ZERO
        ]

        base_fx = Decimal('1')
        if deal_state['agreement_currency'] != 'UZS':
            base_fx = Decimal(str(
                deal_state.get('implied_fx')
                or deal_state.get('purchase_fx')
                or DEFAULT_DEMO_USD_UZS_RATE
            ))
            upsert_exchange_rate(
                tenant_id=ctx['tenant_id'],
                base_currency=deal_state['agreement_currency'],
                quote_currency='UZS',
                rate_date=date.date(),
                rate=base_fx,
                source=ExchangeRate.Source.MANUAL,
                is_manual=True,
                notes=f'Sherik replay rollover FX {deal.deal_id}',
                overwrite_manual=False,
            )

        if investor_topup > ZERO:
            investor_uzs_needed = money(investor_topup * base_fx)
            eligibility = build_payout_decision_preview(
                agreement=deal_state['agreement'],
                decision_type=PayoutDecision.DecisionType.ROLL_OVER_CAPITAL,
                amount_uzs=None,
                allocations=None,
                from_account_id=None,
                now=date,
            )
            eligible_uzs = money(eligibility.get('eligible_total_uzs') or ZERO)
            if eligible_uzs <= ZERO:
                raise CommandError(
                    f'{deal.deal_id}: receipt at {date.isoformat()} needs {fmt(investor_topup)} '
                    f'{deal_state["agreement_currency"]} investor-side extra capital, but no '
                    'recovered capital is eligible for E23 rollover yet.'
                )

            investor_uzs_total = min(investor_uzs_needed, eligible_uzs)
            investor_base_total = money(investor_uzs_total / base_fx)
            if investor_topup - investor_base_total > CENT:
                raise CommandError(
                    f'{deal.deal_id}: receipt at {date.isoformat()} needs {fmt(investor_topup)} '
                    f'{deal_state["agreement_currency"]} investor-side extra capital, but only '
                    f'{fmt(investor_base_total)} {deal_state["agreement_currency"]} is eligible '
                    'from recovered sale proceeds at this point.'
                )

            cash_account = self._select_rollover_source_account(
                ctx=ctx,
                deal_id=deal.deal_id,
                amount_uzs=investor_uzs_total,
                date=date,
                fallback_usd_rate=base_fx,
                target_currency=deal_state['purchase_currency'],
            )
            preview = build_payout_decision_preview(
                agreement=deal_state['agreement'],
                decision_type=PayoutDecision.DecisionType.ROLL_OVER_CAPITAL,
                amount_uzs=investor_uzs_total,
                allocations=None,
                from_account_id=cash_account.id,
                now=date,
            )
            if not preview['allowed']:
                raise CommandError(
                    f'{deal.deal_id}: E23 rollover blocked at {date.isoformat()}: '
                    + '; '.join(preview['blocking_reasons'])
                )
            allocations = [
                {
                    'procurement_id': row['procurement_id'],
                    'partner_id': row['partner_id'],
                    'amount_uzs': money(row['amount_uzs']),
                }
                for row in preview['allocations']
            ]

            execute_payout_decision(
                agreement=deal_state['agreement'],
                decision_type=PayoutDecision.DecisionType.ROLL_OVER_CAPITAL,
                amount_uzs=investor_uzs_total,
                allocations=allocations,
                from_account_id=cash_account.id,
                client_request_id=uuid5(
                    NAMESPACE_URL,
                    f'sherik-rollover:{deal.deal_id}:{source_ref}:{date.isoformat()}:{cash_account.currency}:{investor_uzs_total}',
                ),
                notes=f'Sherik replay capital rollover {deal.deal_id}',
                created_by_id=ctx['owner'].id,
                now=date,
            )

            for row in preview['allocations']:
                amount_uzs = money(row['amount_uzs'])
                amount = money(amount_uzs / base_fx)
                partner_id = int(row['partner_id'])
                deal_state['capital_remaining'][partner_id] = money(
                    deal_state['capital_remaining'].get(partner_id, ZERO) + amount
                )
                deal_state['capital_rollovers'].append({
                    'partner_id': partner_id,
                    'amount': amount,
                    'base_currency': deal_state['agreement_currency'],
                    'source_currency': str(cash_account.currency).upper(),
                    'from_cash_account_id': cash_account.id,
                })

        for row in operator_topups:
            amount = money(row['amount'])
            amount_uzs = money(amount * base_fx)
            partner_id = int(row['partner_id'])
            cash_account = self._select_rollover_source_account(
                ctx=ctx,
                deal_id=deal.deal_id,
                amount_uzs=amount_uzs,
                date=date,
                fallback_usd_rate=base_fx,
                target_currency=deal_state['agreement_currency'],
            )
            add_agreement_contribution(
                tenant_id=ctx['tenant_id'],
                agreement_id=deal_state['agreement'].id,
                partner_id=partner_id,
                amount=amount,
                currency=deal_state['agreement_currency'],
                fx_rate=base_fx,
                date=date,
                notes=f'Sherik replay operator capital top-up {deal.deal_id}',
                client_request_id=str(uuid5(
                    NAMESPACE_URL,
                    f'sherik-operator-capital:{deal.deal_id}:{source_ref}:{date.isoformat()}:{partner_id}:{amount}',
                )),
                created_by_id=ctx['owner'].id,
                actor_partner_id=partner_id,
                from_cash_account_id=cash_account.id,
            )
            deal_state['capital_remaining'][partner_id] = money(
                deal_state['capital_remaining'].get(partner_id, ZERO) + amount
            )

    def _sync_replay_capital_remaining(self, deal_state: dict) -> None:
        from apps.partnerships.workspace_common import _agreement_available_by_partner

        agreement = deal_state['agreement']
        currency_code = str(agreement.currency or 'UZS').upper()
        available = _agreement_available_by_partner(agreement)
        deal_state['capital_remaining'] = {
            int(row['partner_id']): money(
                available.get(int(row['partner_id']), {}).get(currency_code, ZERO)
            )
            for row in deal_state['partners']
        }

    def _select_rollover_source_account(
        self,
        *,
        ctx: dict,
        deal_id: str,
        amount_uzs: Decimal,
        date,
        fallback_usd_rate: Decimal,
        target_currency: str,
    ):
        from apps.finance.fx_rates import upsert_exchange_rate
        from apps.finance.models import CashAccount, ExchangeRate
        from apps.finance.services import exchange_currency

        amount_uzs = money(amount_uzs)
        fallback_usd_rate = Decimal(str(fallback_usd_rate or Decimal('1')))
        if fallback_usd_rate <= ZERO:
            fallback_usd_rate = Decimal('1')
        target_currency = str(target_currency or 'UZS').upper()

        def required_native(currency_code: str) -> Decimal:
            if currency_code == 'USD':
                upsert_exchange_rate(
                    tenant_id=ctx['tenant_id'],
                    base_currency='USD',
                    quote_currency='UZS',
                    rate_date=date.date(),
                    rate=fallback_usd_rate,
                    source=ExchangeRate.Source.MANUAL,
                    is_manual=True,
                    notes=f'Sherik replay rollover source FX {deal_id}',
                    overwrite_manual=False,
                )
                return money(amount_uzs / fallback_usd_rate)
            if currency_code == 'UZS':
                return amount_uzs
            raise CommandError(f'{deal_id}: unsupported rollover source currency {currency_code}.')

        target_account = None
        for account in ctx['cash_accounts'].values():
            account.refresh_from_db()
            if not account.is_active or account.kind != CashAccount.Kind.CASH:
                continue
            currency_code = str(account.currency or 'UZS').upper()
            if currency_code == target_currency:
                target_account = account
                break

        if target_account is None:
            raise CommandError(f'{deal_id}: rollover target cash account {target_currency} was not found.')

        required = required_native(target_currency)
        if money(target_account.balance) + CENT >= required:
            return target_account

        if target_currency == 'USD':
            uzs_account = next(
                (
                    account for account in ctx['cash_accounts'].values()
                    if str(account.currency or '').upper() == 'UZS'
                    and account.kind == CashAccount.Kind.CASH
                    and account.is_active
                ),
                None,
            )
            if uzs_account is not None:
                uzs_account.refresh_from_db()
                target_account.refresh_from_db()
                deficit = money(required - money(target_account.balance))
                inverse_rate = (Decimal('1') / fallback_usd_rate).quantize(Decimal('0.000001'))
                from_amount = money(deficit / inverse_rate)
                while money(from_amount * inverse_rate) + CENT < deficit:
                    from_amount = money(from_amount + CENT)
                if money(uzs_account.balance) + CENT >= from_amount:
                    exchange_currency(
                        tenant_id=ctx['tenant_id'],
                        from_account_id=uzs_account.id,
                        to_account_id=target_account.id,
                        from_amount=from_amount,
                        rate=inverse_rate,
                        date=date,
                        notes=f'Sherik replay implied FX before E23 rollover {deal_id}',
                    )
                    target_account.refresh_from_db()
                    if money(target_account.balance) + CENT >= required:
                        return target_account

        balances = ', '.join(
            f'{account.name}={money(account.balance)} {account.currency}'
            for account in ctx['cash_accounts'].values()
        )
        raise CommandError(
            f'{deal_id}: no operating cash account can fund E23 rollover '
            f'{fmt(amount_uzs)} UZS into {target_currency}. Balances: {balances}.'
        )

    def _transfer_row(self, ctx: dict, state: dict, row: dict) -> None:
        from apps.inventory.models import LotStock
        from apps.inventory.services import transfer_lot_stock

        product = product_key(pick(row, 'MAHSULOT'))
        variant = ctx['variants'].get(product)
        if variant is None:
            raise CommandError(f'transfer product was not created: {product}')
        from_key = norm(pick(row, 'CHIQIM')) or 'ASOSIY'
        to_key = norm(pick(row, 'KIRIM')) or 'DOKON'
        if from_key not in ctx['warehouses'] or to_key not in ctx['warehouses']:
            raise CommandError(f'transfer warehouse is unknown for {product}: {from_key}->{to_key}')
        remaining = int(dec(pick(row, 'SONI')))
        if remaining <= 0:
            return
        self._ensure_stock_at_location(
            ctx=ctx,
            variant=variant,
            destination=ctx['warehouses'][from_key],
            quantity=remaining,
            product=product,
            source=f"transfer row {row.get('_row_id')}",
        )
        stocks = (
            LotStock.objects
            .filter(
                tenant_id=ctx['tenant_id'],
                warehouse=ctx['warehouses'][from_key],
                lot__product_variant=variant,
                quantity_remaining__gt=0,
            )
            .select_related('lot')
            .order_by('lot__received_at', 'lot__id')
        )
        for stock in stocks:
            if remaining <= 0:
                break
            take = min(int(stock.quantity_remaining), remaining)
            transfer_lot_stock(
                tenant_id=ctx['tenant_id'],
                lot=stock.lot,
                from_warehouse=ctx['warehouses'][from_key],
                to_warehouse=ctx['warehouses'][to_key],
                quantity=take,
            )
            remaining -= take
        if remaining > 0:
            raise CommandError(f"Cannot transfer {product}: shortage {remaining} on row {row.get('_row_id')}.")

    def _sale_row(self, ctx: dict, state: dict, row: dict) -> None:
        from apps.customers.models import Customer
        from apps.finance.models import CashAccount
        from apps.sales.models import SalePayment
        from apps.sales.services import create_sale, open_pos_session

        product = product_key(pick(row, 'MAHSULOT'))
        variant = ctx['variants'].get(product)
        if variant is None:
            raise CommandError(f"sale product was not created: {product} row {row.get('_row_id')}.")
        qty_decimal = dec(pick(row, 'JAMI DONA'))
        qty = int(qty_decimal)
        if Decimal(qty) != qty_decimal:
            raise CommandError(f"sale row {row.get('_row_id')} has non-integer quantity {qty_decimal}.")
        sold_at = parse_dt(pick(row, 'SOTUV SANASI'))
        warehouse_key = self._sale_warehouse_key(row)
        location = ctx['warehouses'][warehouse_key]
        session_key = (ctx['tenant_id'], location.id)
        session = state['sessions'].get(session_key)
        if session is None:
            session = open_pos_session(
                tenant_id=ctx['tenant_id'],
                location_id=location.id,
                opened_by_id=ctx['owner'].id,
                opening_cash=ZERO,
                opening_cash_by_currency={'UZS': '0.00', 'USD': '0.00'},
            )
            state['sessions'][session_key] = session
        customer_name = display_name(pick(row, 'MIJOZ')) or 'CUSTOMER'
        customer = Customer.objects.get_or_create(
            tenant=ctx['business'],
            name=customer_name,
            defaults={'is_active': True},
        )[0]
        row_currency = currency(pick(row, 'VALYUTA'))
        unit_raw = dec(pick(row, 'SOTUV NARXI'))
        fx = dec(pick(row, 'KURS'), Decimal('1')) if row_currency == 'USD' else Decimal('1')
        unit_uzs = money(unit_raw if row_currency == 'UZS' else unit_raw * fx)
        method = SalePayment.Method.CASH
        account = (
            CashAccount.objects
            .filter(
                tenant=ctx['business'],
                currency=row_currency,
                kind=CashAccount.Kind.CASH,
                is_active=True,
            )
            .order_by('id')
            .first()
        )
        if account is None:
            account = ctx['cash_accounts']['KASSA SOM']
            row_currency = 'UZS'
        self._ensure_stock_at_location(
            ctx=ctx,
            variant=variant,
            destination=location,
            quantity=qty,
            product=product,
            source=f"sale row {row.get('_row_id')}",
        )
        create_sale(
            tenant_id=ctx['tenant_id'],
            pos_session_id=session.id,
            location_id=location.id,
            sold_by_id=ctx['owner'].id,
            customer_id=customer.id,
            lines=[{
                'product_variant_id': variant.id,
                'quantity': qty,
                'unit_price': unit_uzs,
                'operation_currency': row_currency,
                'operation_unit_price': unit_raw if row_currency == 'USD' else unit_uzs,
                'fx_rate': fx,
                'discount_reason_id': ctx['discount_reason'].id,
            }],
            payments=[{
                'amount': money(unit_raw * qty if row_currency == 'USD' else unit_uzs * qty),
                'currency': row_currency,
                'fx_rate': fx,
                'method': method,
                'account_id': account.id,
            }],
            date=sold_at,
            notes=f"Sherik replay SOTUV row {row.get('_row_id')}",
        )

    def _ensure_stock_at_location(self, *, ctx: dict, variant, destination, quantity: int, product: str, source: str) -> None:
        from django.db.models import Sum

        from apps.inventory.models import LotStock
        from apps.inventory.services import transfer_lot_stock

        current = (
            LotStock.objects
            .filter(
                tenant_id=ctx['tenant_id'],
                warehouse=destination,
                lot__product_variant=variant,
            )
            .aggregate(total=Sum('quantity_remaining'))['total']
            or 0
        )
        remaining = int(quantity) - int(current)
        if remaining <= 0:
            return

        stocks = (
            LotStock.objects
            .filter(
                tenant_id=ctx['tenant_id'],
                lot__product_variant=variant,
                quantity_remaining__gt=0,
            )
            .exclude(warehouse=destination)
            .select_related('lot', 'warehouse')
            .order_by('lot__received_at', 'lot__id')
        )
        for stock in stocks:
            if remaining <= 0:
                break
            take = min(int(stock.quantity_remaining), remaining)
            transfer_lot_stock(
                tenant_id=ctx['tenant_id'],
                lot=stock.lot,
                from_warehouse=stock.warehouse,
                to_warehouse=destination,
                quantity=take,
            )
            remaining -= take
        if remaining > 0:
            raise CommandError(
                f'Cannot reconstruct stock transfer for {product} before {source}: shortage {remaining}.'
            )

    def _close_sessions(self, state: dict) -> None:
        from collections import defaultdict as dd

        from apps.sales.models import Sale, SalePayment
        from apps.sales.services import close_pos_session

        for session in state['sessions'].values():
            cash_totals = dd(Decimal)
            for payment in SalePayment.objects.filter(
                sale__pos_session=session,
                sale__status=Sale.SaleStatus.COMPLETED,
                role=SalePayment.Role.INCOMING,
            ):
                cash_totals[str(payment.currency or 'UZS').upper()] += Decimal(str(payment.amount))
            close_pos_session(
                session=session,
                closed_by_id=session.opened_by_id,
                actual_cash=ZERO,
                actual_cash_by_currency={
                    cur: str(money(amount))
                    for cur, amount in sorted(cash_totals.items())
                },
            )

    def _agreement_partners_payload(self, ctx: dict, agreement_report: dict, agreement_currency: str) -> tuple[list[dict], list[dict]]:
        investor_term = next((term for term in agreement_report['terms'] if 'USTOZ' in norm(term.name)), agreement_report['terms'][0])
        operator_term = next((term for term in agreement_report['terms'] if term is not investor_term), agreement_report['terms'][-1])
        partners = []
        contributions = []
        investor_profit_sum = ZERO
        for split in agreement_report.get('funding_split') or []:
            investor = clean_person_name(split['investor'])
            partner = ctx['investor_partners'].get(investor)
            if partner is None:
                raise CommandError(f"{agreement_report['deal_id']}: missing investor relation for {investor}.")
            amount = split['master_amount_usd'] if agreement_currency == 'USD' else split['contract_capital']
            profit_share = split['profit_share']
            investor_profit_sum += profit_share
            partners.append({
                'partner_id': partner.id,
                'role': 'INVESTOR',
                'planned_capital_share': str(money(amount)),
                'profit_share': str(profit_share),
            })
        initial_splits = agreement_report.get('initial_funding_split')
        if initial_splits is None:
            initial_splits = agreement_report.get('funding_split') or []
        for split in initial_splits:
            investor = clean_person_name(split['investor'])
            partner = ctx['investor_partners'].get(investor)
            if partner is None:
                raise CommandError(f"{agreement_report['deal_id']}: missing investor relation for {investor}.")
            amount = self._funding_amount_for_currency(split, agreement_currency)
            contributions.append({'partner_id': partner.id, 'amount': money(amount)})

        operator_amount = operator_term.capital
        if agreement_currency == 'USD' and agreement_report['terms_currency'] != 'USD':
            implied_fx = agreement_report.get('implied_fx') or Decimal('1')
            operator_amount = money(operator_term.capital / implied_fx) if implied_fx else ZERO
        partners.append({
            'partner_id': ctx['operator'].id,
            'role': 'OPERATOR',
            'planned_capital_share': str(money(operator_amount)),
            'profit_share': str(ratio(Decimal('1') - investor_profit_sum)),
        })
        contributions.append({'partner_id': ctx['operator'].id, 'amount': money(operator_amount)})
        return partners, contributions

    def _capital_allocations_for_amount(
        self,
        partners: list[dict],
        required: Decimal,
        remaining_by_partner: dict[int, Decimal],
    ) -> list[dict]:
        required = money(required)
        def available_amount(partner_id: int) -> Decimal:
            available = money(remaining_by_partner.get(partner_id, ZERO))
            return ZERO if available <= CENT else available

        total = sum((available_amount(int(row['partner_id'])) for row in partners), ZERO)
        if total <= ZERO:
            raise CommandError('Cannot allocate capital without planned capital.')
        if required - total > CENT:
            raise CommandError(f'Cannot allocate {fmt(required)} from remaining capital {fmt(total)}.')

        result = []
        allocated = ZERO
        for index, partner in enumerate(partners):
            partner_id = int(partner['partner_id'])
            available = available_amount(partner_id)
            amount = money(required * available / total)
            if index == len(partners) - 1:
                amount = money(required - allocated)
            if amount - available > CENT:
                amount = available
            allocated += amount
            if amount > ZERO:
                result.append({
                    'partner_id': partner_id,
                    'amount': str(amount),
                })

        residue = money(required - sum((money(row['amount']) for row in result), ZERO))
        if residue > ZERO:
            for partner in partners:
                partner_id = int(partner['partner_id'])
                current = next((row for row in result if int(row['partner_id']) == partner_id), None)
                current_amount = money(current['amount']) if current else ZERO
                headroom = money(available_amount(partner_id) - current_amount)
                if headroom <= ZERO:
                    continue
                top_up = min(headroom, residue)
                if current:
                    current['amount'] = str(money(current_amount + top_up))
                else:
                    result.append({'partner_id': partner_id, 'amount': str(top_up)})
                residue = money(residue - top_up)
                if residue <= ZERO:
                    break

        total_allocated = money(sum((money(row['amount']) for row in result), ZERO))
        if abs(total_allocated - required) > CENT:
            raise CommandError(f'Capital allocation total {fmt(total_allocated)} does not match required {fmt(required)}.')

        for row in result:
            partner_id = int(row['partner_id'])
            amount = money(row['amount'])
            remaining = money(remaining_by_partner.get(partner_id, ZERO) - amount)
            remaining_by_partner[partner_id] = ZERO if abs(remaining) <= CENT else remaining

        return result

    def _capital_topups_for_amount(self, partners: list[dict], amount: Decimal) -> list[dict]:
        amount = money(amount)
        planned_total = sum((money(row['planned_capital_share']) for row in partners), ZERO)
        if amount <= ZERO:
            return []
        if planned_total <= ZERO:
            raise CommandError('Cannot top up capital without planned shares.')

        result = []
        allocated = ZERO
        for index, partner in enumerate(partners):
            partner_id = int(partner['partner_id'])
            planned = money(partner['planned_capital_share'])
            topup = money(amount * planned / planned_total)
            if index == len(partners) - 1:
                topup = money(amount - allocated)
            allocated += topup
            if topup > ZERO:
                result.append({
                    'partner_id': partner_id,
                    'amount': topup,
                })
        return result

    def _purchase_total(self, rows: list[dict]) -> Decimal:
        total = ZERO
        for row in rows:
            total += dec(pick(row, 'SONI')) * dec(pick(row, 'MAHSULOT NARXI'))
        return money(total)

    def _purchase_fx(self, row: dict, purchase_currency: str) -> Decimal:
        return dec(pick(row, 'KURS'), Decimal('1')) if purchase_currency == 'USD' else Decimal('1')

    def _purchase_received_at(self, row: dict):
        override = row.get('_effective_received_at')
        if override:
            return parse_dt(override)
        candidates = [
            parse_dt(value)
            for value in (
                pick(row, 'SANA'),
                pick(row, 'MAHSULOT OMBORGA YETIB KELGAN SANA'),
                pick(row, "MAHSULOT DO'KONGA YETIB KELGAN SANA"),
            )
            if not is_blank(value)
        ]
        return min(candidates) if candidates else timezone.now()

    def _purchase_warehouse_key(self, row: dict) -> str:
        if not is_blank(pick(row, "MAHSULOT DO'KONGA YETIB KELGAN SANA")):
            return 'DOKON'
        return 'ASOSIY'

    def _sale_warehouse_key(self, row: dict) -> str:
        raw = norm(pick(row, 'OMBOR')) or 'DOKON'
        return raw if raw == 'DOKON' else 'DOKON'

    def _row_required_base(self, row: dict, deal_state: dict) -> Decimal:
        row_total = money(dec(pick(row, 'SONI')) * dec(pick(row, 'MAHSULOT NARXI')))
        if deal_state['purchase_currency'] == deal_state['agreement_currency']:
            return row_total
        implied_fx = deal_state.get('implied_fx')
        if not implied_fx:
            raise CommandError('Missing implied FX for cross-currency receive.')
        return money(row_total / implied_fx)

    def _deal_start_at(self, master: dict, deal_id: str, deal: LocalDeal):
        dates = [self._purchase_received_at(row) for row in deal.rows['SOTIB OLISH']]
        raw = (master.get('deals') or {}).get(deal_id, {}).get('Boshlangan sana')
        if raw:
            dates.append(parse_dt(raw))
        return min(dates) if dates else timezone.now()

    def _unique_username(self, base: str, used: set[str]) -> str:
        candidate = base or 'user'
        suffix = 2
        while candidate in used or User.objects.filter(username=candidate).exists():
            candidate = f'{base}_{suffix}'
            suffix += 1
        used.add(candidate)
        return candidate

    def _build_report(self, master: dict, local_deals: dict[str, LocalDeal]) -> dict:
        blockers: list[str] = []
        warnings: list[str] = []
        master_deals = set(master['deals'])
        local_ids = set(local_deals)
        for correction in master.get('source_corrections') or []:
            warnings.append(f'master: source correction applied: {correction}.')
        for deal_id in sorted(master_deals - local_ids):
            warnings.append(f'{deal_id}: master deal has no local workbook; replay will skip it.')
        for deal_id in sorted(local_ids - master_deals):
            warnings.append(f'{deal_id}: local workbook has no master BITIMLAR row.')

        agreements = []
        products_by_business = defaultdict(lambda: defaultdict(set))
        payout_reports = []
        stock_reports = []

        for deal_id, deal in sorted(local_deals.items()):
            for correction in deal.source_corrections:
                warnings.append(f'{deal_id}: source correction applied: {correction}.')
            agreement = self._agreement_report(deal_id, deal, master, blockers, warnings)
            agreements.append(agreement)
            purchase_total = self._purchase_total(deal.rows['SOTIB OLISH'])
            if purchase_total - agreement.get('total_capital', ZERO) > CENT:
                warnings.append(
                    f"{deal_id}: purchase total {fmt(purchase_total)} {deal.terms_currency} exceeds "
                    f"initial agreement capital {fmt(agreement['total_capital'])}; apply will require "
                    'E23 capital rollover facts from recovered sale proceeds.'
                )
            payout_reports.append(self._payout_report(deal_id, deal, master, blockers, warnings))
            stock_reports.append(self._stock_report(deal_id, deal, blockers, warnings))

            operator_name = agreement.get('operator_name') or ''
            business = business_key(deal_id, operator_name)
            for row in deal.rows['SOTIB OLISH']:
                key = product_key(pick(row, 'MAHSULOT'))
                if key:
                    products_by_business[business][key].add(deal_id)

        timeline_stock = self._timeline_stock_report(local_deals, agreements, blockers, warnings)

        product_reports = []
        for business, product_map in sorted(products_by_business.items()):
            repeated = {
                name: sorted(deals)
                for name, deals in product_map.items()
                if len(deals) > 1
            }
            product_reports.append({
                'business': business,
                'unique_products': len(product_map),
                'repeated_products': repeated,
            })

        return {
            'blockers': blockers,
            'warnings': warnings,
            'master_deals': sorted(master_deals),
            'local_deals': sorted(local_ids),
            'agreements': agreements,
            'payouts': payout_reports,
            'stock': stock_reports,
            'timeline_stock': timeline_stock,
            'products': product_reports,
        }

    def _funding_split_from_rows(
        self,
        rows: list[dict],
        investor_term: AgreementTerm,
        master_pool: Decimal,
    ) -> list[dict]:
        if not rows or investor_term.capital <= ZERO or master_pool <= ZERO:
            return []
        funding_amounts = defaultdict(Decimal)
        funding_dates = {}
        row_ids = defaultdict(list)
        source_dates = {}
        funding_sources = {}
        for row in rows:
            investor = row['investor']
            funding_amounts[investor] += row['amount_usd']
            funding_dates.setdefault(investor, row['date'])
            source_dates.setdefault(investor, row.get('source_date') or row['date'])
            funding_sources.setdefault(investor, row.get('funding_source') or '')
            row_ids[investor].append(row.get('row_id'))

        allocated = ZERO
        close_to_total = abs(sum(funding_amounts.values(), ZERO) - master_pool) <= CENT
        result = []
        sorted_rows = [
            {'investor': investor, 'amount_usd': amount, 'date': funding_dates.get(investor)}
            for investor, amount in sorted(funding_amounts.items())
        ]
        for index, row in enumerate(sorted_rows):
            pool_share = ratio(row['amount_usd'] / master_pool)
            contract_capital = money(investor_term.capital * pool_share)
            if close_to_total and index == len(sorted_rows) - 1:
                contract_capital = money(investor_term.capital - allocated)
            allocated += contract_capital
            investor = row['investor']
            result.append({
                'investor': investor,
                'master_amount_usd': row['amount_usd'],
                'first_date': row['date'],
                'source_date': source_dates.get(investor),
                'funding_source': funding_sources.get(investor, ''),
                'row_ids': [row_id for row_id in row_ids.get(investor, []) if row_id],
                'pool_share': pool_share,
                'contract_capital': contract_capital,
                'profit_share': ratio(investor_term.profit_share * pool_share),
            })
        return result

    def _funding_row_payload(
        self,
        row: dict,
        investor_term: AgreementTerm,
        master_pool: Decimal,
    ) -> dict:
        pool_share = ratio(row['amount_usd'] / master_pool) if master_pool > ZERO else ZERO
        return {
            'investor': row['investor'],
            'master_amount_usd': row['amount_usd'],
            'date': row['date'],
            'source_date': row.get('source_date') or row['date'],
            'funding_source': row.get('funding_source') or '',
            'row_id': row.get('row_id'),
            'pool_share': pool_share,
            'contract_capital': money(investor_term.capital * pool_share),
            'profit_share': ratio(investor_term.profit_share * pool_share),
        }

    def _funding_amount_for_currency(self, row: dict, agreement_currency: str) -> Decimal:
        if str(agreement_currency or '').upper() == 'USD':
            return money(row['master_amount_usd'])
        return money(row['contract_capital'])

    def _agreement_report(self, deal_id: str, deal: LocalDeal, master: dict, blockers: list[str], warnings: list[str]) -> dict:
        if len(deal.terms) < 2:
            blockers.append(f'{deal_id}: FOYDA_TAQSIMOTI must contain investor and operator rows.')
            return {'deal_id': deal_id, 'ok': False, 'partners': []}

        investor_term = next((term for term in deal.terms if 'USTOZ' in norm(term.name)), deal.terms[0])
        operator_term = next((term for term in deal.terms if term is not investor_term), deal.terms[-1])
        total_capital = sum((term.capital for term in deal.terms), ZERO)
        master_funding = master['funding_by_deal'].get(deal_id, [])
        master_pool = sum((row['amount_usd'] for row in master_funding), ZERO)

        formula_ok = False
        mudaraba_ratio = ZERO
        operator_expected = ZERO
        if investor_term.capital_share > ZERO:
            mudaraba_ratio = ratio(investor_term.profit_share / investor_term.capital_share)
            operator_expected = ratio(
                operator_term.capital_share
                + ((Decimal('1') - mudaraba_ratio) * investor_term.capital_share)
            )
            formula_ok = (
                Decimal('0') <= mudaraba_ratio <= Decimal('1')
                and abs(operator_expected - operator_term.profit_share) <= Decimal('0.0001')
            )
        if not formula_ok:
            blockers.append(f'{deal_id}: FOYDA_TAQSIMOTI profit/capital shares do not fit current agreement formula.')

        profit_sum = sum((term.profit_share for term in deal.terms), ZERO)
        if abs(profit_sum - Decimal('1')) > Decimal('0.000001'):
            blockers.append(f'{deal_id}: partner profit shares sum to {fmt(profit_sum)}, expected 1.')

        local_pool = investor_term.capital
        pool_status = 'ok'
        implied_fx = None
        if master_pool:
            if deal.terms_currency == 'USD':
                diff = abs(local_pool - master_pool)
                if diff > Decimal('0.11'):
                    blockers.append(f'{deal_id}: local investor pool {fmt(local_pool)} USD differs from master {fmt(master_pool)} USD.')
                    pool_status = 'mismatch'
                elif diff:
                    pool_status = f'rounding {fmt(diff)} USD'
            else:
                implied_fx = ratio(local_pool / master_pool) if master_pool > ZERO else None
                pool_status = f'fx {fmt(implied_fx)} {deal.terms_currency}/USD'

        funding_split = []
        initial_funding_split = []
        supplemental_funding = []
        if master_funding and investor_term.capital > ZERO and master_pool > ZERO:
            deal_start = self._deal_start_at(master, deal_id, deal)
            earliest_funding_at = min(parse_dt(row['date']) for row in master_funding)
            initial_rows = []
            supplemental_rows = []
            for row in master_funding:
                row_at = parse_dt(row['date'])
                if row_at <= deal_start or row_at == earliest_funding_at:
                    initial_rows.append(row)
                else:
                    supplemental_rows.append(row)

            funding_split = self._funding_split_from_rows(master_funding, investor_term, master_pool)
            initial_funding_split = self._funding_split_from_rows(initial_rows, investor_term, master_pool)
            supplemental_funding = [
                self._funding_row_payload(row, investor_term, master_pool)
                for row in supplemental_rows
            ]

        return {
            'deal_id': deal_id,
            'ok': formula_ok,
            'terms_currency': deal.terms_currency,
            'total_capital': total_capital,
            'investor_pool': local_pool,
            'master_pool': master_pool,
            'pool_status': pool_status,
            'implied_fx': implied_fx,
            'mudaraba_ratio': mudaraba_ratio,
            'operator_expected': operator_expected,
            'operator_name': operator_term.name,
            'terms': deal.terms,
            'funding_split': funding_split,
            'initial_funding_split': initial_funding_split,
            'supplemental_funding': supplemental_funding,
        }

    def _payout_report(self, deal_id: str, deal: LocalDeal, master: dict, blockers: list[str], warnings: list[str]) -> dict:
        sale_cash = defaultdict(Decimal)
        capital_cash = defaultdict(Decimal)
        for row in deal.rows['TUSHUM']:
            account, account_currency, amount = account_amount(row)
            row_type = norm(pick(row, 'TUSHUM TURI'))
            if row_type == 'SOTUV':
                sale_cash[(account, account_currency)] += amount
            elif row_type == 'CAPITAL':
                capital_cash[(account, account_currency)] += amount

        fx_in = defaultdict(Decimal)
        fx_out = defaultdict(Decimal)
        for row in deal.rows['PUL AYRIBOSHLASH']:
            in_account = norm(pick(row, 'KIRIM'))
            out_account = norm(pick(row, 'CHIQIM'))
            out_currency = currency(pick(row, 'VALYUTA (CHIQUVCHI)'))
            in_currency = 'USD' if in_account == 'KASSA DOLLAR' else 'UZS' if in_account in {'KASSA SOM', 'PLASTIK SOM'} else ''
            fx_in[(in_account, in_currency)] += money(pick(row, 'KIRIM MIQDORI'))
            fx_out[(out_account, out_currency)] += money(pick(row, 'CHIQIM MIQDORI'))

        payouts = master['payouts_by_deal'].get(deal_id, [])
        withdrawals = master['withdrawals_by_deal'].get(deal_id, [])
        payouts_by_currency = defaultdict(Decimal)
        capital_return = ZERO
        profit_payout = ZERO
        for row in payouts:
            payouts_by_currency[row['currency']] += row['amount']
            if 'KAPITAL' in norm(row['kind']):
                capital_return += row['amount']
            elif 'FOYDA' in norm(row['kind']):
                profit_payout += row['amount']

        released_usd = sum((row['amount'] for row in withdrawals), ZERO)
        payout_total = sum(payouts_by_currency.values(), ZERO)
        usd_sale_plus_fx = (
            sale_cash.get(('KASSA DOLLAR', 'USD'), ZERO)
            + fx_in.get(('KASSA DOLLAR', 'USD'), ZERO)
        )
        if payout_total and released_usd + CENT < payout_total:
            blockers.append(f'{deal_id}: master released USD {fmt(released_usd)} is less than investor payouts {fmt(payout_total)}.')
        if payouts_by_currency.get('USD', ZERO) and usd_sale_plus_fx + CENT < payouts_by_currency['USD']:
            warnings.append(
                f'{deal_id}: explicit USD sale/FX cash {fmt(usd_sale_plus_fx)} is less than USD payouts '
                f'{fmt(payouts_by_currency["USD"])}; may need final-date FX or delayed payout.'
            )

        return {
            'deal_id': deal_id,
            'tushum_rows': len(deal.rows['TUSHUM']),
            'fx_rows': len(deal.rows['PUL AYRIBOSHLASH']),
            'sale_cash': dict(sale_cash),
            'capital_cash': dict(capital_cash),
            'fx_in': dict(fx_in),
            'fx_out': dict(fx_out),
            'released_usd': released_usd,
            'payout_total': payout_total,
            'capital_return': capital_return,
            'profit_payout': profit_payout,
            'payouts_by_currency': dict(payouts_by_currency),
            'usd_sale_plus_fx': usd_sale_plus_fx,
        }

    def _stock_report(self, deal_id: str, deal: LocalDeal, blockers: list[str], warnings: list[str]) -> dict:
        purchased = defaultdict(Decimal)
        sold = defaultdict(Decimal)
        sale_posted = Counter()
        for row in deal.rows['SOTIB OLISH']:
            key = product_key(pick(row, 'MAHSULOT'))
            if key:
                purchased[key] += dec(pick(row, 'SONI'))
        for row in deal.rows['SOTUV']:
            key = product_key(pick(row, 'MAHSULOT'))
            qty = dec(pick(row, 'JAMI DONA'))
            if not key or qty <= ZERO:
                continue
            sold[key] += qty
            sale_posted[str(pick(row, 'POSTED') or '')] += 1

        missing = []
        for key, sold_qty in sorted(sold.items()):
            bought_qty = purchased.get(key, ZERO)
            if sold_qty > bought_qty:
                missing.append({
                    'product': key,
                    'sold': sold_qty,
                    'purchased': bought_qty,
                    'shortage': sold_qty - bought_qty,
                })
        for row in missing:
            blockers.append(
                f"{deal_id}: sales exceed purchased quantity for {row['product']}: "
                f"sold {fmt(row['sold'])}, purchased {fmt(row['purchased'])}, shortage {fmt(row['shortage'])}."
            )

        if sale_posted and any(key in {'', 'None'} for key in sale_posted):
            warnings.append(
                f'{deal_id}: SOTUV POSTED is incomplete/inconsistent; replay treats required SOTUV rows as source facts.'
            )

        return {
            'deal_id': deal_id,
            'products_sold': len(sold),
            'missing': missing,
            'sale_posted': dict(sale_posted),
        }

    def _timeline_stock_report(
        self,
        local_deals: dict[str, LocalDeal],
        agreements: list[dict],
        blockers: list[str],
        warnings: list[str],
    ) -> list[dict]:
        agreement_by_deal = {row['deal_id']: row for row in agreements}
        events = []
        for deal_id, deal in sorted(local_deals.items()):
            agreement = agreement_by_deal.get(deal_id, {})
            business = business_key(deal_id, clean_person_name(agreement.get('operator_name') or ''))
            for row in deal.rows['SOTIB OLISH']:
                key = product_key(pick(row, 'MAHSULOT'))
                qty = dec(pick(row, 'SONI'))
                if key and qty > ZERO:
                    events.append({
                        'business': business,
                        'at': self._purchase_received_at(row),
                        'order': 20,
                        'deal_id': deal_id,
                        'row_id': int(row.get('_row_id') or 0),
                        'kind': 'receive',
                        'product': key,
                        'warehouse': self._purchase_warehouse_key(row),
                        'qty': qty,
                    })
            for row in deal.rows['STOCK TRANSFER']:
                key = product_key(pick(row, 'MAHSULOT'))
                qty = dec(pick(row, 'SONI'))
                if key and qty > ZERO:
                    events.append({
                        'business': business,
                        'at': parse_dt(pick(row, 'SANA')),
                        'order': 30,
                        'deal_id': deal_id,
                        'row_id': int(row.get('_row_id') or 0),
                        'kind': 'transfer',
                        'product': key,
                        'from_warehouse': norm(pick(row, 'CHIQIM')) or 'ASOSIY',
                        'to_warehouse': norm(pick(row, 'KIRIM')) or 'DOKON',
                        'qty': qty,
                    })
            for row in deal.rows['SOTUV']:
                key = product_key(pick(row, 'MAHSULOT'))
                qty = dec(pick(row, 'JAMI DONA'))
                if key and qty > ZERO:
                    events.append({
                        'business': business,
                        'at': parse_dt(pick(row, 'SOTUV SANASI')),
                        'order': 40,
                        'deal_id': deal_id,
                        'row_id': int(row.get('_row_id') or 0),
                        'kind': 'sale',
                        'product': key,
                        'warehouse': self._sale_warehouse_key(row),
                        'qty': qty,
                    })

        balances = defaultdict(Decimal)
        hard_shortages = []
        reconstructed_transfers = []

        def move_available(event: dict, *, destination: str, needed: Decimal) -> Decimal:
            remaining = needed
            for (business, warehouse, product), available in sorted(list(balances.items())):
                if business != event['business'] or product != event['product'] or warehouse == destination:
                    continue
                if available <= ZERO:
                    continue
                take = min(available, remaining)
                balances[(business, warehouse, product)] -= take
                balances[(business, destination, product)] += take
                remaining -= take
                reconstructed_transfers.append({
                    **event,
                    'from_warehouse': warehouse,
                    'to_warehouse': destination,
                    'qty': take,
                })
                if remaining <= ZERO:
                    break
            return remaining

        # Separate missing transfer evidence from true business-level shortages.
        for event in sorted(events, key=lambda item: (
            item['business'],
            item['at'],
            item['order'],
            item['deal_id'],
            item['row_id'],
        )):
            business = event['business']
            product = event['product']
            if event['kind'] == 'receive':
                balances[(business, event['warehouse'], product)] += event['qty']
                continue

            if event['kind'] == 'transfer':
                src = (business, event['from_warehouse'], product)
                dst = (business, event['to_warehouse'], product)
                available = balances[src]
                if available < event['qty']:
                    remaining = move_available(event, destination=event['from_warehouse'], needed=event['qty'] - available)
                    available = balances[src]
                    if remaining > ZERO or available < event['qty']:
                        hard_shortages.append({
                            **event,
                            'warehouse': event['from_warehouse'],
                            'available': available,
                            'shortage': event['qty'] - available,
                        })
                balances[src] -= event['qty']
                balances[dst] += event['qty']
                continue

            if event['kind'] == 'sale':
                key = (business, event['warehouse'], product)
                available = balances[key]
                if available < event['qty']:
                    remaining = move_available(event, destination=event['warehouse'], needed=event['qty'] - available)
                    available = balances[key]
                    if remaining > ZERO or available < event['qty']:
                        hard_shortages.append({
                            **event,
                            'available': available,
                            'shortage': event['qty'] - available,
                        })
                balances[key] -= event['qty']

        summarized = {}
        for shortage in hard_shortages:
            key = (
                shortage['business'],
                shortage['warehouse'],
                shortage['product'],
            )
            current = summarized.get(key)
            if current is None or shortage['shortage'] > current['shortage']:
                summarized[key] = shortage

        shortage_report = []
        for shortage in sorted(summarized.values(), key=lambda item: (
            item['business'],
            item['product'],
            item['warehouse'],
        )):
            blockers.append(
                f"{shortage['business']}: timeline stock shortage for {shortage['product']} "
                f"at {shortage['warehouse']} on {shortage['at'].date().isoformat()} "
                f"({shortage['kind']} {shortage['deal_id']} row {shortage['row_id']}): "
                f"available {fmt(shortage['available'])}, need {fmt(shortage['qty'])}."
            )
            shortage_report.append({
                'business': shortage['business'],
                'product': shortage['product'],
                'warehouse': shortage['warehouse'],
                'at': shortage['at'],
                'kind': shortage['kind'],
                'deal_id': shortage['deal_id'],
                'row_id': shortage['row_id'],
                'available': shortage['available'],
                'qty': shortage['qty'],
                'shortage': shortage['shortage'],
            })

        if reconstructed_transfers:
            warnings.append(
                f'Timeline stock preflight reconstructed {len(reconstructed_transfers)} missing stock transfer(s); '
                'apply will need explicit transfer_lot_stock events before affected sales.'
            )

        return {
            'shortages': shortage_report,
            'reconstructed_transfers': reconstructed_transfers,
        }

    def _print_report(self, report: dict, *, stage: str) -> None:
        self.stdout.write('Sherik Excel Replay preflight')
        self.stdout.write(f"  master deals: {', '.join(report['master_deals'])}")
        self.stdout.write(f"  local deals: {', '.join(report['local_deals'])}")

        if stage in {'plan', 'preflight'}:
            self.stdout.write('\nAgreements')
            for row in report['agreements']:
                self.stdout.write(
                    f"  {row['deal_id']}: total={fmt(row['total_capital'])} {row['terms_currency']}, "
                    f"investor_pool={fmt(row['investor_pool'])} {row['terms_currency']}, "
                    f"master_pool={fmt(row['master_pool'])} USD, "
                    f"mudaraba_ratio={fmt(row['mudaraba_ratio'])}, pool={row['pool_status']}"
                )
                for term in row['terms']:
                    self.stdout.write(
                        f"    {term.name}: capital={fmt(term.capital)}, "
                        f"capital_share={fmt(term.capital_share)}, profit_share={fmt(term.profit_share)}"
                    )
                for split in row['funding_split']:
                    self.stdout.write(
                        f"    funding {split['investor']}: master={fmt(split['master_amount_usd'])} USD, "
                        f"pool_share={fmt(split['pool_share'])}, contract_capital={fmt(split['contract_capital'])}, "
                        f"profit_share={fmt(split['profit_share'])}"
                    )
                for split in row.get('supplemental_funding') or []:
                    self.stdout.write(
                        f"    supplemental funding {split['investor']}: "
                        f"effective={date_label(split['date'])}, source={date_label(split['source_date'])}, "
                        f"amount={fmt(split['master_amount_usd'])} USD"
                    )

        if stage == 'preflight':
            self.stdout.write('\nPayout and FX coverage')
            for row in report['payouts']:
                self.stdout.write(
                    f"  {row['deal_id']}: released={fmt(row['released_usd'])} USD, "
                    f"payouts={fmt(row['payout_total'])} USD "
                    f"(capital={fmt(row['capital_return'])}, profit={fmt(row['profit_payout'])}), "
                    f"usd_sale_plus_fx={fmt(row['usd_sale_plus_fx'])}, "
                    f"tushum_rows={row['tushum_rows']}, fx_rows={row['fx_rows']}"
                )

            self.stdout.write('\nStock coverage')
            for row in report['stock']:
                missing = ', '.join(
                    f"{item['product']} shortage={fmt(item['shortage'])}"
                    for item in row['missing']
                )
                self.stdout.write(
                    f"  {row['deal_id']}: products_sold={row['products_sold']}, "
                    f"missing={missing or '-'}"
                )

            self.stdout.write('\nTimeline stock coverage')
            if report['timeline_stock']['shortages']:
                for row in report['timeline_stock']['shortages']:
                    self.stdout.write(
                        f"  {row['business']}: {row['product']} at {row['warehouse']} "
                        f"{row['at'].date().isoformat()} {row['kind']} "
                        f"{row['deal_id']} row {row['row_id']}, "
                        f"available={fmt(row['available'])}, need={fmt(row['qty'])}"
                    )
            else:
                self.stdout.write('  ok')
            if report['timeline_stock']['reconstructed_transfers']:
                self.stdout.write(
                    f"  reconstructed_transfers={len(report['timeline_stock']['reconstructed_transfers'])}"
                )

            self.stdout.write('\nProduct catalog by business')
            for row in report['products']:
                repeated = ', '.join(
                    f"{name} ({'/'.join(deals)})"
                    for name, deals in sorted(row['repeated_products'].items())
                )
                self.stdout.write(
                    f"  {row['business']}: unique_products={row['unique_products']}, "
                    f"repeated={repeated or '-'}"
                )

        if report['warnings']:
            self.stdout.write('\nWarnings')
            for warning in report['warnings']:
                self.stdout.write(f'  WARN: {warning}')

        if report['blockers']:
            self.stdout.write('\nBlockers')
            for blocker in report['blockers']:
                self.stdout.write(f'  BLOCKER: {blocker}')
        else:
            self.stdout.write('\nBlockers: none')
