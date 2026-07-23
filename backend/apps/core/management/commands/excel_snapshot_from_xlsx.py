"""Build canonical Excel-alignment snapshot JSON from a local .xlsx file."""

from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

from django.core.management.base import BaseCommand, CommandError

from openpyxl import load_workbook


DEFAULT_SHEETS = (
    'MALUMOTLAR',
    'SOTIB OLISH',
    'STOCK TRANSFER',
    'SOTUV',
    'TUSHUM',
    'XARAJAT',
    'PUL AYRIBOSHLASH',
)


def _norm_key(value: str) -> str:
    return ' '.join(str(value).replace('\n', ' ').strip().upper().split())


def _pick(row: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        target = _norm_key(key)
        for row_key, row_value in row.items():
            if _norm_key(row_key) == target:
                return row_value
    return ''


def _is_blank(value: Any) -> bool:
    return value is None or str(value).strip() == ''


def _serialize(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return value


def _is_required_row(sheet_name: str, row: dict[str, Any]) -> bool:
    if sheet_name == 'MALUMOTLAR':
        return any(
            not _is_blank(_pick(row, key))
            for key in ('TRADE CREDITOR', 'MIJOZLAR', 'MAHSULOTLAR', 'OMBOR')
        )
    if sheet_name == 'SOTIB OLISH':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SANA', 'MAHSULOT', 'SONI', 'MAHSULOT NARXI')
        )
    if sheet_name == 'STOCK TRANSFER':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SANA', 'MAHSULOT', 'SONI', 'KIRIM', 'CHIQIM')
        )
    if sheet_name == 'SOTUV':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SOTUV SANASI', 'MAHSULOT', 'JAMI DONA', 'SOTUV NARXI')
        )
    if sheet_name == 'TUSHUM':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SANA', 'MIJOZ', 'MIQDOR')
        )
    if sheet_name == 'XARAJAT':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SANA', 'MIQDOR')
        )
    if sheet_name == 'PUL AYRIBOSHLASH':
        return all(
            not _is_blank(_pick(row, key))
            for key in ('SANA', 'KIRIM', 'CHIQIM', 'CHIQIM MIQDORI')
        )
    return False


class Command(BaseCommand):
    help = (
        'Parse a local .xlsx and build canonical snapshot JSON for '
        'excel_workflow_staged / excel_workflow_audit.'
    )

    def add_arguments(self, parser):
        parser.add_argument(
            '--source-xlsx',
            type=str,
            required=True,
            help='Path to source .xlsx file.',
        )
        parser.add_argument(
            '--output-json',
            type=str,
            required=True,
            help='Path where snapshot JSON will be saved.',
        )
        parser.add_argument(
            '--include-sheets',
            nargs='*',
            default=list(DEFAULT_SHEETS),
            help='Optional subset of sheet names to include.',
        )

    def handle(self, *args, **options):
        source_path = Path(options['source_xlsx']).expanduser()
        output_path = Path(options['output_json']).expanduser()
        include_sheets = tuple(options.get('include_sheets') or DEFAULT_SHEETS)

        if not source_path.exists():
            raise CommandError(f'Source file not found: {source_path}')

        workbook = load_workbook(source_path, data_only=True)
        missing = [name for name in include_sheets if name not in workbook.sheetnames]
        if missing:
            raise CommandError(f'Missing sheets in workbook: {missing}')

        snapshot: dict[str, Any] = {
            'sheets': {},
            'opening_balances': {},
            'expected': {},
        }

        for sheet_name in include_sheets:
            ws = workbook[sheet_name]
            headers = [
                str(cell.value).strip() if cell.value is not None else ''
                for cell in ws[1]
            ]
            sheet_rows: list[dict[str, Any]] = []

            for row_id, row_cells in enumerate(
                ws.iter_rows(min_row=2, max_row=ws.max_row),
                start=2,
            ):
                row: dict[str, Any] = {}
                for header, cell in zip(headers, row_cells):
                    if not header:
                        continue
                    row[header] = _serialize(cell.value)

                if all(_is_blank(value) for value in row.values()):
                    continue
                if not _is_required_row(sheet_name, row):
                    continue

                row['_row_id'] = str(row_id)
                sheet_rows.append(row)

            snapshot['sheets'][sheet_name] = sheet_rows
            self.stdout.write(f'{sheet_name}: {len(sheet_rows)} rows')

        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(
            json.dumps(snapshot, ensure_ascii=False, indent=2),
            encoding='utf-8',
        )
        self.stdout.write(self.style.SUCCESS(f'Snapshot saved: {output_path}'))
